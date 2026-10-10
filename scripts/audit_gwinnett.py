"""Replay Gwinnett County's 2024–25 source/cohort evidence without fitting or importing."""
import argparse
from collections import Counter, defaultdict
from datetime import date, datetime
import csv
import hashlib
import io
import json
from pathlib import Path
import re
import statistics
import subprocess
import zipfile

import prepare_ccd_states as georgia
from state_registry import reported_count

ROOT = Path(__file__).resolve().parents[1]
OUTPUT = ROOT / 'data/source/gwinnett-district-audit.json'
LEA = '1302550'
NATIVE_LEA = 'GA-667'
YEAR = '2024-2025'
OPERATIONAL = {'1', '3', '4', '5', '8'}
GRADES = ['PK', 'KG'] + [str(g) for g in range(1, 14)] + ['UG', 'AE']
LOWER_GRADES = ['Pre-Kindergarten', 'Kindergarten'] + ['Grade ' + str(g) for g in range(1, 9)]
STATE_EXTRACT_SHA = '3b6e55bba0055c582eb14a461697c64131a36e4cd555bca94965e5f77cac91fc'
GEORGIA_SNAPSHOT_SHA = '46d47899274835d484be76ac9b13515afb96474c54b5a676c4c43539df289a67'
RAW_INPUT_SHA = '5f6e99cede2abfc989372132f8b3a58327eaa68a104d443d7855479a258f0242'
SOURCES = {'ccd_directory_2025': {'path': 'data/raw/ccd_sch_029_2425_w_1a_073025.zip', 'url': 'https://nces.ed.gov/ccd/Data/zip/ccd_sch_029_2425_w_1a_073025.zip', 'bytes': 13352819, 'sha256': '39326da788aa322353d20ceaf8ad4baed26272502cd05b066cf6c594988b21ab'}, 'ccd_membership_2025': {'path': 'data/raw/ccd_sch_052_2425_l_1a_073025.zip', 'url': 'https://nces.ed.gov/ccd/Data/zip/ccd_sch_052_2425_l_1a_073025.zip', 'bytes': 212696691, 'sha256': '4a7f660c5fc5eaae488dd02fd43498f349fc828b227edd0970d5b6995ead4d4d'}, 'ccd_lunch_2025': {'path': 'data/raw/ccd_sch_033_2425_l_2a_073025.zip', 'url': 'https://nces.ed.gov/ccd/Data/zip/ccd_sch_033_2425_l_2a_073025.zip', 'bytes': 14029439, 'sha256': '97bda749e778ee74cb731d181bbcdaf0f9c6cd6edf94cb31518a5bbeab411dd6'}, 'ccd_state_notes_2025': {'path': 'data/raw/SY_2024-25_CCD_Final_2a_Data_Notes.xlsx', 'url': 'https://nces.ed.gov/ccd/xls/SY_2024-25_CCD_Final_2a_Data_Notes.xlsx', 'bytes': 210984, 'sha256': '6476b6b28b700389e261b87377d9cd249935503d53b5919fe1b7ef8b1052d1a7'}, 'edc_codebook_v3_1': {'path': 'data/raw/EDC_codebook_v3.1.xlsx', 'url': 'https://www.eddatacenter.org/data_codebooks/EDC_codebook_v3.1.xlsx', 'bytes': 20401, 'sha256': 'cfaf3eeb482b416106e4a7b775c9554768fbf4900b0518c861979d646b6b2574'}, 'edc_technical_v3_1': {'path': 'data/raw/EDC_technical_documentation_v3.1.pdf', 'url': 'https://www.eddatacenter.org/data_documentation/EDC_technical_documentation_v3.1.pdf', 'bytes': 1100602, 'sha256': '749c52497e37cfd73c0e29afadea3d90d6b50950844f070247594f44254a7968'}, 'nces_lunch_definitions': {'path': 'data/raw/ccd-lunch-definition.html', 'url': 'https://nces.ed.gov/learn/blog/understanding-school-lunch-eligibility-common-core-data', 'bytes': 34948, 'sha256': '03f24547726ebf7529d946cee85f84e3de3ac3962b5dbc8b96e0a4db3744247d'}, 'edc_2025': {'path': 'data/raw/edc-2025/ga.csv', 'url': 'https://www.eddatacenter.org/api/data/3.1?state=GA&year=2025', 'bytes': 93271612, 'sha256': '38e3c39bdebd361d55055478e2f657a9e8c70d60781b9e60f47407c955f8cc6d'}}

POLICY = {
    'roster': 'Exact operational 2024–25 CCD LEA 1302550 / native GA-667 attachments. Retain complete directory, membership, lunch and native All Students Math/ELA records, including all high/mixed observations; names never join.',
    'identity': 'Join authoritative 12-digit NCES ID and year. EDC StateAssignedDistID is 667 and StateAssignedSchID already contains 667-ssss; cross-check GA- + that exact native school code against CCD ST_SCHID. No name matching or numeric school-code padding.',
    'offered_vs_enrolled': 'Retain complete offered flags and native enrolled subtotals separately. Strict pure lower offerings additionally exclude ungraded/adult/high grades. Source eligibility reuses the audited Georgia complete lower-grade and reconciled membership contract, never discovery counts.',
    'grade_school_contract': 'As reported complete Yes/No offers, highest grade 01–08, no offered grades 9–13/adult, reported PK–8 subtotals sum to total, all outside grades explicitly Reported/Derived zero, and explicitly reported-zero Ungraded if offered. Primary, zero-membership and unreported configurations remain distinct.',
    'assessment': 'Only exact published School/All Students/All Students Georgia Milestones EOG Regular G38 Math/ELA, Levels 3-4, 2024–25, EDC V3.1 may enter source cohorts. SEA-native aggregate provenance supports point rates; source business rules and valid-score denominators remain unverified. Never average grade percentages or extend high/mixed populations.',
    'income': georgia.DEFINITION,
    'income_metadata': 'Direct certification is a same-year benefits-eligibility proxy, not full FRPL or all economic disadvantage. Retain Georgia CCD notes on Medicaid-income use starting 2024 alongside SNAP/TANF and foster records, and fall state-funded membership including outside-program PK. Neither source describes exactly the tested G38 children.',
    'charters_types_virtual': 'Retain exact charter, alternative and other type flags. CCD has no virtual-status field; native EDC SchVirtual remains separately labeled. Labels and fit quality do not filter source membership. Historical state-note program IDs never create current directory records.',
    'counts': 'EDC tested/proficient counts and participation remain raw unverified evidence. Every verified denominator, sampling variance and interval remains null/unavailable; CCD enrollment never substitutes or weights averaged grades.',
    'combined': 'Equal mean of eligible native Math and ELA proficiency, not the fraction proficient in both.',
    'approval': 'Source/cohort audit only: status and scope audit_pending, source/modeling approval false. Separate district fits/diagnostics, canonical import and browser verification remain required.',
}
PLANNING = {'source': 'data/source/district-comparison-candidates.json', 'operational_roster': 141,
            'enrollment': 182518, 'potential_ES': 111, 'potential_HS': 24,
            'explanation': 'Discovery only. Reconcile exact school IDs and original membership separately; planning counts do not select outcomes, income or model members and do not approve high-school assessments.'}


def native_cell(value):
    # Workbook date cells remain explicitly typed ISO values in JSON; no count,
    # text, boolean, zero, missing or suppression cell is coerced.
    if isinstance(value, (datetime, date)):
        return {'cell_type': type(value).__name__, 'iso8601': value.isoformat()}
    return value


def fingerprint(value):
    return hashlib.sha256(json.dumps(value, sort_keys=True, separators=(',', ':'),
                                    ensure_ascii=False).encode()).hexdigest()


def load_georgia(path=georgia.EXTRACT):
    content = Path(path).read_bytes()
    if hashlib.sha256(content).hexdigest() != STATE_EXTRACT_SHA:
        raise ValueError('Committed state extract fingerprint changed; fresh district source audit required')
    payload = json.loads(content)
    snapshot = next(s for s in payload['states'] if s['state']['id'] == 'GA')
    validate_georgia(snapshot)
    return snapshot


def validate_georgia(snapshot):
    if fingerprint(snapshot) != GEORGIA_SNAPSHOT_SHA or fingerprint(snapshot['sources']) != fingerprint(SOURCES):
        raise ValueError('Georgia native source contract/fingerprint changed')
    # Validation only; never calls prepare_snapshot, SQLite or fitting.
    georgia.validate_snapshot(snapshot)


def ccd_identity(raw):
    if (raw.get('SCHOOL_YEAR') != YEAR or raw.get('ST') != 'GA'
            or raw.get('LEAID') != LEA or raw.get('ST_LEAID') != NATIVE_LEA
            or not re.fullmatch(LEA + r'\d{5}', raw.get('NCESSCH', ''))
            or not re.fullmatch(r'GA-667-\d{4}', raw.get('ST_SCHID', ''))):
        raise ValueError('Gwinnett source requires exact same-year CCD school/LEA identifiers')
    return raw['NCESSCH']


def edc_identity(raw, directory):
    key = georgia.nces_id(raw.get('NCESSchoolID', ''), 'GA')
    district, school = raw.get('StateAssignedDistID', ''), raw.get('StateAssignedSchID', '')
    if (district != '667' or not re.fullmatch(r'667-\d{4}', school)
            or raw.get('NCESDistrictID') != LEA or district != '667'
            or raw.get('StateAbbrev') != 'GA' or raw.get('StateFips') != '13'
            or raw.get('SchYear') != '2024-25' or raw.get('Version') != 'V3.1'
            or raw.get('DataLevel') != 'School' or raw.get('StudentGroup') != 'All Students'
            or raw.get('StudentSubGroup') != 'All Students' or raw.get('Subject') not in {'math', 'ela'}
            or raw.get('GradeLevel') not in {'G03', 'G04', 'G05', 'G06', 'G07', 'G08', 'G38'}
            or raw.get('AssmtName') != 'Georgia Milestones EOG' or raw.get('AssmtType') != 'Regular'
            or raw.get('ProficiencyCriteria') != 'Levels 3-4' or not key.startswith(LEA)):
        raise ValueError('EDC record is outside the exact Georgia native assessment/identity contract')
    native = 'GA-' + school
    if key in directory and directory[key]['raw']['ST_SCHID'] != native:
        raise ValueError('EDC native code disagrees with exact NCES/CCD school identity')
    return key


def offered_evidence(raw):
    flags = {g: raw.get('G_' + g + '_OFFERED') for g in GRADES}
    complete = (raw.get('IGOFFERED') == 'As reported' and raw.get('NOGRADES') == 'No'
                and all(v in {'Yes', 'No'} for v in flags.values()))
    offered = [g for g, v in flags.items() if v == 'Yes']
    lower, high = any(g in offered for g in GRADES[:10]), any(str(g) in offered for g in range(9, 13))
    scope = ('uncertain' if not complete else 'other' if '13' in offered or 'AE' in offered
             or not (lower or high) else 'mixed' if lower and high else 'lower' if lower else 'high')
    return {'flags': flags, 'complete_reported_flags': complete, 'graded_scope': scope,
            'ungraded_offered': flags['UG'], 'strict_pure_offered_lower': scope == 'lower' and flags['UG'] == 'No',
            'offered': offered, 'offered_tested_grades': [str(g) for g in range(3, 9) if str(g) in offered],
            'low': raw.get('GSLO'), 'high': raw.get('GSHI'), 'level': raw.get('LEVEL'),
            'scope_note': 'Graded scope disregards UG for description only; native eligibility additionally requires reported-zero UG and reconciled lower membership.'}


def enrolled_evidence(raw, total, grade_entries, offered):
    grades = {g: e['raw'] for g, e in grade_entries.items()}
    values = {g: {'source_row': e['source_row'], 'raw_count': e['raw']['STUDENT_COUNT'],
                  'dms_flag': e['raw']['DMS_FLAG'], 'reported_count': reported_count(e['raw'])}
              for g, e in sorted(grade_entries.items())}
    lower = [reported_count(r) for g, r in grades.items() if g in LOWER_GRADES]
    lower_sum = sum(lower) if lower and all(v is not None for v in lower) else None
    total_count = reported_count(total)
    native = georgia.grade_school(raw, grades, total)
    evidence = [g for g, v in values.items() if v['reported_count'] is not None and v['reported_count'] > 0]
    offered_native = ['Pre-Kindergarten' if g == 'PK' else 'Kindergarten' if g == 'KG'
                      else 'Ungraded' if g == 'UG' else 'Adult Education' if g == 'AE' else 'Grade ' + g
                      for g in offered['offered']]
    return {'native_grade_school_contract': native, 'grade_values': values,
            'reported_lower_subtotal': lower_sum, 'reported_total': total_count,
            'lower_total_reconciled': lower_sum is not None and lower_sum == total_count,
            'ungraded_reported_count': reported_count(grades.get('Ungraded')),
            'outside_lower_explicit_zero': bool(grades) and all(
                r['STUDENT_COUNT'] == '0' and r['DMS_FLAG'] in {'Reported', 'Derived'}
                for g, r in grades.items() if g not in LOWER_GRADES),
            'positive_enrolled_grades': evidence,
            'positive_enrolled_without_offer': [g for g in evidence if g not in offered_native],
            'offered_explicit_zero_enrolled': [g for g in offered_native if g in values and values[g]['reported_count'] == 0],
            'missing_offered_enrollment': [g for g in offered_native if g not in values],
            'unreported_enrolled_grades': [g for g, v in values.items() if v['reported_count'] is None]}


def income_evidence(raw, total_entry, direct_entry, grade_entries):
    total, direct = (None if total_entry is None else total_entry['raw']), (None if direct_entry is None else direct_entry['raw'])
    n, dc = reported_count(total), reported_count(direct)
    income_reason = ('Same-year CCD membership is missing, nonreported or zero' if n is None or n <= 0
                     else 'CCD direct certification is absent' if direct is None
                     else 'CCD direct certification is suppressed or nonreported' if dc is None
                     else 'Direct-certification count exceeds individual-school membership' if dc > n else None)
    native_reason = georgia.eligibility('GA', raw, total, direct, {g: e['raw'] for g, e in grade_entries.items()})
    return {'source_id': 'ccd_lunch_2025', 'year': 2025,
            'membership_source_id': 'ccd_membership_2025',
            'membership_source_row': None if total_entry is None else total_entry['source_row'],
            'direct_source_row': None if direct_entry is None else direct_entry['source_row'],
            'reported_membership': n, 'reported_direct_certification': dc,
            'raw_membership': None if total is None else total['STUDENT_COUNT'],
            'membership_flag': None if total is None else total['DMS_FLAG'],
            'raw_direct_certification': None if direct is None else direct['STUDENT_COUNT'],
            'direct_flag': None if direct is None else direct['DMS_FLAG'],
            'direct_status': 'absent' if direct is None else direct['DMS_FLAG'],
            'percentage': None if income_reason else 100 * dc / n,
            'income_exclusion': income_reason, 'native_source_contract_exclusion': native_reason}


def subject_evidence(entries, configuration, applicable, income):
    values = {}
    for subject, native in [('math', 'math'), ('reading', 'ela')]:
        entry = entries.get(native)
        raw = None if entry is None else entry['raw']
        rate = None if raw is None else georgia.exact_rate(raw['ProficientOrAbove_percent'])
        outcome_reason = ('Missing native school All Students G38 subject row' if entry is None
                          else 'Proficiency suppressed, ranged or missing' if rate is None else None)
        scope_reason = ('Not an unambiguous grade school in the same-year CCD directory' if not configuration
                        else 'No offered grade 3–8 in native lower configuration' if not applicable else None)
        reason = scope_reason or income['income_exclusion'] or outcome_reason
        values[subject] = {'source_id': 'edc_2025', 'source_row': None if entry is None else entry['source_row'],
                           'compact_source_row': None if entry is None else entry['compact_source_row'],
                           'applicable': applicable, 'native_rate': None if rate is None else 100 * rate,
                           'raw_rate': None if raw is None else raw['ProficientOrAbove_percent'],
                           'raw_tested': None if raw is None else raw['StudentSubGroup_TotalTested'],
                           'raw_proficient': None if raw is None else raw['ProficientOrAbove_count'],
                           'raw_participation': None if raw is None else raw['ParticipationRate'],
                           'outcome_exclusion': outcome_reason, 'scope_exclusion': scope_reason,
                           'income_exclusion': income['income_exclusion'], 'exclusion': reason,
                           'usable': reason is None, 'valid_scores': None,
                           'sampling_variance': None, 'sampling_interval_available': False}
    usable = all(v['usable'] for v in values.values())
    reason = '; '.join(s + ': ' + v['exclusion'] for s, v in values.items() if v['exclusion']) or None
    values['combined'] = {'applicable': applicable, 'usable': usable,
                          'native_rate': (values['math']['native_rate'] + values['reading']['native_rate']) / 2 if usable else None,
                          'exclusion': reason, 'valid_scores': None, 'sampling_variance': None,
                          'sampling_interval_available': False, 'definition': POLICY['combined']}
    return values


def summarize(links, predicate, basis):
    members = [r for r in links if predicate(r)]
    subjects = {}
    for subject in ['math', 'reading', 'combined']:
        eligible = [r for r in members if r['subjects'][subject]['usable']]
        subjects[subject] = {'applicable': sum(r['subjects'][subject]['applicable'] for r in members),
                             'usable': len(eligible), 'prospective_30_school_floor_met': len(eligible) >= 30, 'eligible_school_ids': [r['school_id'] for r in eligible],
                             'exclusions': dict(sorted(Counter(r['subjects'][subject]['exclusion'] for r in members if not r['subjects'][subject]['usable']).items())),
                             'verified_valid_score_counts': 0, 'model_sampling_intervals_available': False}
    eligible = [r for r in members if r['subjects']['combined']['usable']]
    incomes = [r['income']['percentage'] for r in eligible]
    return {'profiles': len(members), 'school_ids': [r['school_id'] for r in members],
            'membership_basis': basis, 'applicable_basis': 'Native grade-school contract plus at least one offered grade 3–8; does not assert every offered grade is tested or enrollment equals valid scores.',
            'subjects': subjects,
            'combined_ccd_charter': dict(sorted(Counter(r['ccd_flags']['charter'] for r in eligible).items())),
            'combined_ccd_school_type': dict(sorted(Counter(r['ccd_flags']['school_type'] for r in eligible).items())),
            'combined_income': {'minimum': min(incomes) if incomes else None,
                                'maximum': max(incomes) if incomes else None, 'distinct': len(set(incomes)),
                                'mean': statistics.mean(incomes) if incomes else None,
                                'sample_standard_deviation': statistics.stdev(incomes) if len(incomes) > 1 else None}}


def derive(document, snapshot):
    if (type(document.get('schema_version')) is not int or document.get('schema_version') != 1 or document.get('state') != 'GA'
            or type(document.get('year')) is not int or document.get('year') != 2025 or document.get('academic_year') != YEAR
            or document.get('nces_lea_id') != LEA or document.get('native_lea_id') != NATIVE_LEA
            or document.get('status') != 'audit_pending' or document.get('scope') != 'audit_pending'
            or document.get('approved_for_source') is not False or document.get('approved_for_modeling') is not False
            or fingerprint(document.get('sources')) != fingerprint(SOURCES) or fingerprint(document.get('policy')) != fingerprint(POLICY)
            or fingerprint(document.get('planning')) != fingerprint(PLANNING)
            or document.get('state_extract') != {'path': 'data/source/ccd-state-snapshots.json',
                'sha256': STATE_EXTRACT_SHA, 'georgia_snapshot_sha256': GEORGIA_SNAPSHOT_SHA}):
        raise ValueError('Gwinnett source-only audit contract/provenance changed')
    validate_georgia(snapshot)
    inputs = document['raw_inputs']
    if fingerprint(inputs) != RAW_INPUT_SHA:
        raise ValueError('Retained complete Gwinnett raw records/headers/row references changed')
    directory, directory_rows, totals, grade_maps, direct, by_school = {}, set(), {}, defaultdict(dict), {}, defaultdict(dict)
    for entry in inputs['directory']['records']:
        key = ccd_identity(entry['raw'])
        if key in directory or entry['source_row'] in directory_rows or type(entry['source_row']) is not int or entry['source_row'] < 2:
            raise ValueError('Duplicate CCD identity or directory source row')
        directory[key] = entry
        directory_rows.add(entry['source_row'])
    for source_key in ['membership', 'lunch']:
        seen = set()
        for entry in inputs[source_key]['records']:
            raw = entry['raw']; key = ccd_identity(raw)
            if entry['source_row'] in seen or type(entry['source_row']) is not int or entry['source_row'] < 2:
                raise ValueError('Duplicate or invalid CCD raw source row')
            seen.add(entry['source_row'])
            if key in directory and raw['ST_SCHID'] != directory[key]['raw']['ST_SCHID']:
                raise ValueError('CCD native school ID disagrees across same-year sources')
            if source_key == 'membership':
                if raw['TOTAL_INDICATOR'] == 'Education Unit Total':
                    if key in totals: raise ValueError('Duplicate school membership total')
                    totals[key] = entry
                elif raw['TOTAL_INDICATOR'] == 'Subtotal 4 - By Grade':
                    if raw['GRADE'] in grade_maps[key]: raise ValueError('Duplicate enrolled grade subtotal')
                    grade_maps[key][raw['GRADE']] = entry
                else: raise ValueError('Unexpected retained membership aggregation')
            elif raw['DATA_GROUP'] == 'Direct Certification':
                if key in direct or raw['TOTAL_INDICATOR'] != 'Education Unit Total':
                    raise ValueError('Duplicate or non-school Direct Certification')
                direct[key] = entry
    native_seen, compact_seen = set(), set()
    for entry in inputs['assessment']['records']:
        raw = entry['raw']; key = edc_identity(raw, directory)
        unique = (key, raw['GradeLevel'], raw['Subject'])
        if unique in native_seen or entry['compact_source_row'] in compact_seen or type(entry['source_row']) is not int or entry['source_row'] < 2:
            raise ValueError('Duplicate native school/grade/subject or invalid row reference')
        native_seen.add(unique); compact_seen.add(entry['compact_source_row'])
        if raw['GradeLevel'] == 'G38': by_school[key][raw['Subject']] = entry
    from plan_district_comparisons import potential_cohort
    discovery = inputs['planning']['record']
    if (discovery['nces_lea_id'] != LEA or discovery['native_lea_id'] != NATIVE_LEA
            or discovery['state'] != 'GA' or type(discovery['enrollment']) is not int
            or discovery['enrollment'] != 182518 or discovery['enrollment_flag'] != 'Reported'
            or discovery['school_directory_operational_count'] != 141):
        raise ValueError('Exact same-year discovery agency identity/enrollment changed')
    potential = {level: sorted(key for key, entry in directory.items()
        if entry['raw']['UPDATED_STATUS'] in OPERATIONAL and potential_cohort(entry['raw'],
            None if key not in totals else totals[key]['raw'], {g: e['raw'] for g, e in grade_maps[key].items()}) == level)
        for level in ['ES', 'HS']}
    if (fingerprint(potential) != fingerprint(discovery['potential_school_ids'])
            or fingerprint({k: len(v) for k, v in potential.items()}) != fingerprint(discovery['potential_cohorts'])):
        raise ValueError('Planning school identities differ from exact original membership reconstruction')
    profiles = {r['id']: r for r in snapshot['schools'] if r['district_id'] == LEA}
    links, nonoperational = [], []
    for key, entry in directory.items():
        raw = entry['raw']; offered = offered_evidence(raw)
        total = None if key not in totals else totals[key]['raw']
        enrolled = enrolled_evidence(raw, total, grade_maps[key], offered)
        configuration = enrolled['native_grade_school_contract']
        applicable = configuration and bool(offered['offered_tested_grades'])
        income = income_evidence(raw, totals.get(key), direct.get(key), grade_maps[key])
        native = by_school[key]
        subjects = subject_evidence(native, configuration, applicable, income)
        pair = set(native) == {'math', 'ela'}
        if native and not pair: raise ValueError('Incomplete native G38 subject pair')
        source_profile = configuration and income['income_exclusion'] is None and pair
        link = {'school_id': key, 'native_school_id': raw['ST_SCHID'], 'name': raw['SCH_NAME'],
                'ccd_source_row': entry['source_row'], 'offered': offered, 'enrolled': enrolled,
                'income': income, 'subjects': subjects, 'native_g38_pair': pair,
                'native_g38_source_rows': {s: e['source_row'] for s, e in sorted(native.items())},
                'edc_metadata': {s: {k: e['raw'][k] for k in ['SchName', 'DistName', 'SchType', 'SchLevel', 'SchVirtual', 'DistCharter', 'DistType', 'Flag_AssmtNameChange', 'Flag_CutScoreChange_ELA', 'Flag_CutScoreChange_math']} for s, e in sorted(native.items())},
                'ccd_flags': {'operational': raw['UPDATED_STATUS'] in OPERATIONAL, 'status': raw['UPDATED_STATUS'],
                    'status_text': raw['UPDATED_STATUS_TEXT'], 'school_type_id': raw['SCH_TYPE'],
                    'school_type': raw['SCH_TYPE_TEXT'], 'charter': raw['CHARTER_TEXT'],
                    'virtual_status': None, 'virtual_status_note': 'No virtual-status field in this CCD directory; EDC SchVirtual is retained separately.'},
                'historical_state_source_profile': key in profiles,
                'source_profile_eligible': source_profile,
                'state_snapshot_profile': profiles.get(key)}
        if raw['UPDATED_STATUS'] not in OPERATIONAL:
            link['exclusion'] = 'CCD status is nonoperational for the same-year district baseline'
            nonoperational.append(link); continue
        if source_profile != (key in profiles):
            raise ValueError('Raw district source membership disagrees with pinned Georgia source snapshot')
        if key in profiles:
            profile = profiles[key]
            if fingerprint([profile['income'], profile['enrollment'], profile['low_income']]) != fingerprint([income['percentage'], income['reported_membership'], income['reported_direct_certification']]):
                raise ValueError('Raw DC/member arithmetic disagrees with approved native snapshot')
            for k, v in profile['directory_raw'].items():
                if fingerprint(raw[k]) != fingerprint(v): raise ValueError('Raw directory disagrees with approved native snapshot')
            for subject, native_subject in [('math', 'math'), ('reading', 'ela')]:
                outcome, e = profile['outcomes'][subject], native[native_subject]
                if (fingerprint(outcome['raw']) != fingerprint({k: e['raw'][k] for k in georgia.RAW_OUTCOME_FIELDS})
                        or type(outcome['source_row']) is not int or outcome['source_row'] != e['compact_source_row']
                        or outcome['proficiency'] != subjects[subject]['native_rate'] or outcome['tested'] is not None):
                    raise ValueError('Native outcome/source row disagrees with approved snapshot')
        links.append(link)
    op_ids = {r['school_id'] for r in links}
    configuration = lambda r: r['enrolled']['native_grade_school_contract']
    applicable = lambda r: configuration(r) and bool(r['offered']['offered_tested_grades'])
    cohorts = {
        'native_grade_school_configurations': summarize(links, configuration, 'Native complete lower-grade/reconciled membership contract; primary and zero totals remain distinct.'),
        'native_offered_tested_grade': summarize(links, applicable, 'Native configuration with at least one offered grade 3–8.'),
        'native_positive_membership_tested_candidates': summarize(links, lambda r: applicable(r) and r['income']['reported_membership'] is not None and r['income']['reported_membership'] > 0, 'Native applicable configuration with positive reported enrollment; planning cross-check only.'),
        'native_configurations_with_usable_income': summarize(links, lambda r: configuration(r) and r['income']['income_exclusion'] is None, 'Native configuration plus positive reported membership/individual Direct Certification.'),
        'historical_native_source_profiles': summarize(links, lambda r: r['source_profile_eligible'], 'Exact native G38 pair plus native configuration and usable same-year DC income; independently equals pinned state source subset.'),
        'native_primary_only': summarize(links, lambda r: configuration(r) and not r['offered']['offered_tested_grades'], 'Native lower configuration without offered grade 3–8; not assessment-applicable.'),
        'mixed_offers_with_native_g38': summarize(links, lambda r: r['offered']['graded_scope'] == 'mixed' and r['native_g38_pair'], 'Native G38 observations do not approve a mixed-school regression population.'),
    }
    coverage = {
        'directory_records': len(directory), 'operational_schools': len(links), 'nonoperational_schools': len(nonoperational),
        'operational_statuses': dict(sorted(Counter(r['ccd_flags']['status_text'] for r in links).items())),
        'nonoperational_statuses': dict(sorted(Counter(r['ccd_flags']['status_text'] for r in nonoperational).items())),
        'operational_graded_offered_scope': dict(sorted(Counter(r['offered']['graded_scope'] for r in links).items())),
        'operational_ungraded_offered': sum(r['offered']['ungraded_offered'] == 'Yes' for r in links),
        'operational_ccd_school_types': dict(sorted(Counter(r['ccd_flags']['school_type'] for r in links).items())),
        'operational_ccd_charter': dict(sorted(Counter(r['ccd_flags']['charter'] for r in links).items())),
        'membership_totals': len(totals), 'membership_grade_maps': sum(bool(v) for v in grade_maps.values()),
        'all_lunch_rows': len(inputs['lunch']['records']), 'direct_certification_rows': len(direct),
        'direct_certification_flags': dict(sorted(Counter(e['raw']['DMS_FLAG'] for e in direct.values()).items())),
        'operational_usable_income': sum(r['income']['income_exclusion'] is None for r in links),
        'native_configurations_income_exclusions': dict(sorted(Counter(r['income']['income_exclusion'] for r in links if configuration(r) and r['income']['income_exclusion']).items())),
        'assessment_rows': len(inputs['assessment']['records']),
        'assessment_grade_rows': dict(sorted(Counter(e['raw']['GradeLevel'] for e in inputs['assessment']['records']).items())),
        'native_g38_pairs': sum(r['native_g38_pair'] for r in links),
        'g38_edc_virtual_rows': dict(sorted(Counter(e['raw']['SchVirtual'] for e in inputs['assessment']['records'] if e['raw']['GradeLevel'] == 'G38').items())),
        'historical_native_source_profiles': len(profiles),
        'lower_offered_primary_only': sum(r['offered']['graded_scope'] == 'lower' and not r['offered']['offered_tested_grades'] for r in links),
        'lower_offered_outside_native_contract': sum(r['offered']['graded_scope'] == 'lower' and not configuration(r) for r in links),
        'positive_enrolled_without_offer_schools': sum(bool(r['enrolled']['positive_enrolled_without_offer']) for r in links),
        'verified_valid_score_counts': 0, 'sampling_intervals_available': False,
    }
    outside = {source: [e for e in inputs[source]['records'] if (georgia.nces_id(e['raw']['NCESSchoolID'], 'GA') if source == 'assessment' else e['raw']['NCESSCH']) not in op_ids]
               for source in ['membership', 'lunch', 'assessment']}
    return {'roster_links': links, 'nonoperational_directory': nonoperational,
            'outside_operational_roster_records': outside, 'cohorts': cohorts, 'coverage': coverage,
            'historical_state_lunch_notes': snapshot['source_audit']['state_lunch_notes'],
            'later_release_note_roster_gaps': sorted({school_id for sheet in inputs['definitions']['georgia_rows_and_headers']
                if sheet['worksheet'] == 'Data Notes' for entry in sheet['rows']
                for school_id in re.findall(LEA + r'\d{5}', ' '.join(str(v) for v in entry['raw']))} - set(directory)),
            'state_definition_evidence': inputs['definitions'],
            'planning_identity_reconciliation': {'potential_school_ids': potential,
                'source_native_eligible_school_ids': cohorts['historical_native_source_profiles']['subjects']['combined']['eligible_school_ids'],
                'source_not_in_planning_ES': sorted(set(cohorts['historical_native_source_profiles']['subjects']['combined']['eligible_school_ids']) - set(potential['ES'])),
                'planning_ES_not_source_eligible': sorted(set(potential['ES']) - set(cohorts['historical_native_source_profiles']['subjects']['combined']['eligible_school_ids'])),
                'scope_note': 'Discovery reconstructed separately; never a source-income/outcome membership filter.'},
            'historical_state_source_approval': {'approved_point_only': True,
                'basis': snapshot['source_audit']['approval_basis'],
                'district_source_approval': False, 'district_modeling_approval': False}}


def validate(document, snapshot=None):
    expected = derive(document, snapshot if snapshot is not None else load_georgia())
    source_fields = {'schema_version', 'state', 'year', 'academic_year', 'nces_lea_id',
                     'native_lea_id', 'status', 'scope', 'approved_for_source',
                     'approved_for_modeling', 'sources', 'policy', 'planning',
                     'raw_inputs', 'state_extract'}
    if set(document) != source_fields | set(expected):
        raise ValueError('Unexpected or missing source-only audit fields')
    if any(fingerprint(document.get(k)) != fingerprint(v) for k, v in expected.items()):
        raise ValueError('Gwinnett links, raw missingness, exclusions or source cohorts failed offline replay')
    return expected['coverage']


def extract_raw_inputs(membership_evidence=None):
    for metadata in SOURCES.values():
        path = ROOT / metadata['path']
        if hashlib.sha256(path.read_bytes()).hexdigest() != metadata['sha256'] or path.stat().st_size != metadata['bytes']:
            raise ValueError('Pinned raw source fingerprint changed: ' + metadata['path'])
    inputs = {}
    for name, source_id in [('directory', 'ccd_directory_2025'), ('lunch', 'ccd_lunch_2025')]:
        path = ROOT / SOURCES[source_id]['path']
        member = path.stem + '.csv'
        with zipfile.ZipFile(path) as archive, archive.open(member) as stream:
            reader = csv.DictReader(io.TextIOWrapper(stream, encoding='utf-8-sig'))
            records = [{'source_row': i, 'raw': r} for i, r in enumerate(reader, 2)
                       if r['LEAID'] == LEA or r['ST_LEAID'] == NATIVE_LEA]
            header = reader.fieldnames
        inputs[name] = {'source_id': source_id, 'member': member, 'header_source_row': 1,
                        'header': header, 'selection': 'All district records by exact LEA or native LEA; identities cross-checked.', 'records': records}
    source_id = 'ccd_membership_2025'
    if membership_evidence is not None:
        evidence = json.loads(Path(membership_evidence).read_text())
        expected = SOURCES[source_id]
        if any(evidence.get(k) != expected[k] for k in ['path', 'url', 'sha256']):
            raise ValueError('Membership bootstrap provenance disagrees with pinned original archive')
        inputs['membership'] = {'source_id': source_id, 'member': evidence['member'],
            'header_source_row': 1, 'header': evidence['header'],
            'selection': 'Complete exact district Education Unit Total and Subtotal 4 - By Grade from the original archive; all raw fields and CSV row references retained. Race/sex subgroups never become denominators.',
            'records': evidence['records']}
    else:
        # A default explicit extraction always reopens the original Deflate64
        # archive. The committed subset is used only by ordinary offline replay.
        path = ROOT / SOURCES[source_id]['path']; member = path.stem + '.csv'
        process = subprocess.Popen(['unzip', '-p', str(path), member], stdout=subprocess.PIPE,
                                   text=True, encoding='utf-8-sig')
        try:
            reader = csv.DictReader(process.stdout)
            records = [{'source_row': i, 'raw': row} for i, row in enumerate(reader, 2)
                if (row['LEAID'] == LEA or row['ST_LEAID'] == NATIVE_LEA)
                and row['TOTAL_INDICATOR'] in {'Education Unit Total', 'Subtotal 4 - By Grade'}]
            if process.wait() != 0:
                raise ValueError('Failed streaming pinned original school membership archive')
            header = reader.fieldnames
        finally:
            process.stdout.close()
            if process.poll() is None:
                process.terminate(); process.wait()
        inputs['membership'] = {'source_id': source_id, 'member': member,
            'header_source_row': 1, 'header': header,
            'selection': 'Complete exact district Education Unit Total and Subtotal 4 - By Grade from the original archive; all raw fields and CSV row references retained. Race/sex subgroups never become denominators.',
            'records': records}
    import openpyxl
    workbook = openpyxl.load_workbook(ROOT / SOURCES['ccd_state_notes_2025']['path'], read_only=True, data_only=True)
    definitions = []
    try:
        for sheet in workbook:
            rows = [{'source_row': i, 'raw': [native_cell(v) for v in row]} for i, row in enumerate(sheet.iter_rows(values_only=True), 1)
                    if i <= 3 or any(isinstance(v, str) and ('Georgia' in v or 'Gwinnett' in v or v == 'GEORGIA') for v in row)]
            if any(r['source_row'] > 3 for r in rows):
                definitions.append({'worksheet': sheet.title, 'rows': rows})
    finally:
        workbook.close()
    workbook = openpyxl.load_workbook(ROOT / SOURCES['edc_codebook_v3_1']['path'], read_only=True, data_only=True)
    try:
        codebook = [{'worksheet': sheet.title, 'rows': [{'source_row': i, 'raw': [native_cell(v) for v in row]}
                    for i, row in enumerate(sheet.iter_rows(values_only=True), 1)]} for sheet in workbook]
    finally:
        workbook.close()
    inputs['definitions'] = {'ccd_source_id': 'ccd_state_notes_2025', 'georgia_rows_and_headers': definitions,
                             'codebook_source_id': 'edc_codebook_v3_1', 'complete_codebook_worksheets': codebook}
    planning_payload = json.loads((ROOT / PLANNING['source']).read_text())
    agency = next(r for r in planning_payload['agencies'] if r['nces_lea_id'] == LEA)
    fields = ['nces_lea_id', 'state', 'native_lea_id', 'enrollment', 'enrollment_flag',
              'school_directory_operational_count', 'potential_cohorts', 'potential_school_ids',
              'directory_raw', 'directory_source_row', 'membership_raw', 'membership_source_row']
    inputs['planning'] = {'source': PLANNING['source'], 'sources': planning_payload['sources'],
                          'record': {k: agency[k] for k in fields}}
    with (ROOT / SOURCES['edc_2025']['path']).open(encoding='utf-8-sig', newline='') as stream:
        reader = csv.DictReader(stream); records = []; compact_row = 0
        for i, row in enumerate(reader, 2):
            if row['DataLevel'] != 'School' or row['StudentGroup'] != 'All Students' or row['StudentSubGroup'] != 'All Students' or row['Subject'] not in {'math', 'ela'}: continue
            compact_row += 1
            if row['NCESDistrictID'] == LEA or row['StateAssignedDistID'] == '667':
                records.append({'source_row': i, 'compact_source_row': compact_row, 'raw': row})
        header = reader.fieldnames
    inputs['assessment'] = {'source_id': 'edc_2025', 'header_source_row': 1, 'header': header,
                            'selection': 'All exact district School/All Students/All Students Math/ELA grade rows; only native G38 rates can enter source cohorts.',
                            'compact_source_row_basis': 'Ordinal in reproducible statewide Georgia All Students school Math/ELA compact extract, starting at 1; full CSV source_row starts at 2.', 'records': records}
    return inputs


def extract(output=OUTPUT, membership_evidence=None):
    snapshot = load_georgia()
    inputs = extract_raw_inputs(membership_evidence)
    document = {'schema_version': 1, 'state': 'GA', 'year': 2025, 'academic_year': YEAR,
                'nces_lea_id': LEA, 'native_lea_id': NATIVE_LEA,
                'status': 'audit_pending', 'scope': 'audit_pending',
                'approved_for_source': False, 'approved_for_modeling': False,
                'sources': SOURCES, 'policy': POLICY, 'planning': PLANNING, 'raw_inputs': inputs,
                'state_extract': {'path': 'data/source/ccd-state-snapshots.json', 'sha256': STATE_EXTRACT_SHA,
                                  'georgia_snapshot_sha256': GEORGIA_SNAPSHOT_SHA}}
    document.update(derive(document, snapshot))
    validate(document, snapshot)
    Path(output).write_text(json.dumps(document, separators=(',', ':'), ensure_ascii=False, allow_nan=False) + '\n')
    return document['coverage']


if __name__ == '__main__':
    parser = argparse.ArgumentParser(description=__doc__)
    parser.add_argument('--extract', action='store_true')
    parser.add_argument('--membership-evidence', type=Path, help='Bootstrap from a separately verified original district membership subset')
    parser.add_argument('--output', type=Path, default=OUTPUT)
    args = parser.parse_args()
    result = extract(args.output, args.membership_evidence) if args.extract else validate(json.loads(args.output.read_text()))
    print(json.dumps(result, sort_keys=True))
