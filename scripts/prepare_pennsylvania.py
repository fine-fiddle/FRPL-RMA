"""Pennsylvania 2025 native PSSA school totals -> SQLite -> static JSON.

Default rebuilds offline from the committed extract. --extract reads the
official raw files documented in docs/pennsylvania-data.md.
"""
import argparse
from collections import Counter
from datetime import date
import hashlib
import json
import math
from pathlib import Path
import re

import openpyxl
import polars as pl

from database import ROOT, DEFAULT_DB, connect, definition_id, save_models
from prepare_data import build_history

DATASET = 'pa-pssa'
YEAR = 2025
ASSESSMENT = 'PSSA grades 3–8 · grade schools'
EXTRACT = ROOT / 'data/source/pennsylvania.json'
FOLDER = ROOT / 'data/pennsylvania'
BASE = 'https://www.pa.gov/content/dam/copapwp-pagov/en/education/documents/'
PORTAL = 'https://www.pa.gov/agencies/education/data-and-reporting/assessment-reporting'
ENROLLMENT = 'https://www.pa.gov/agencies/education/data-and-reporting/enrollment'
FILES = {
    'pa-assessment-2025': ('pa_assessment_2025.xlsx', BASE+'data-and-reporting/pssa-and-ayp-results/2025-pssa-school-level-data.xlsx'),
    'pa-income-2025': ('pa_income_2025.xlsx', BASE+'data-and-reporting/loan-cancellation/public/2425%20public%20schools%20percent%20low%20income.xlsx'),
    'pa-grades-2025': ('pa_grades_2025.xlsx', BASE+'data-and-reporting/enrollment/public-school/enrollment%20public%20schools%202024-25.xlsx'),
    'pa-pims-2025': ('pa_pims_2025.pdf', BASE+'data-and-reporting/pims/pims-manuals/2024-2025%20pims%20manual%20vol%201.pdf'),
}
ECON_DEFINITION = (
    'Pennsylvania 2024–25 October 1 PIMS individual Economic Disadvantaged Status '
    'Code, field 88: low-income enrollment / total enrollment in the same-year '
    'public-school low-income workbook. Poverty sources can include TANF, census '
    'poor, Medicaid, neglected/delinquent institutions and foster homes; recent '
    'reliable free/reduced-price lunch eligibility is a fallback. CEP schools must '
    'use individual poverty sources, not universal meal eligibility (separate '
    'food-program field 131). Local source choices vary; this is the reported '
    'state-specific economic-disadvantage proxy, not uniform FRPL or direct income. '
    'Denominator includes pre-K/kindergarten where enrolled. Reporting coverage is unknown.')
ASSESSMENT_DEFINITION = (
    'Published 2025 PSSA All Students / Total school rows, grades 3–8, ELA and Math; '
    'Percent Proficient and above and Number Scored. Native school totals are used '
    'without aggregating rounded grade percentages or reconstructing suppression. '
    'PSSA general assessment only, excluding PASA alternate assessment. Source excludes '
    'students enrolled after October 1, homeschoolers, excluded school aggregations, '
    'ELL enrolled for less than one year, LIFE students entering EL after the previous '
    'testing window and students who did not attempt. Rates suppressed below 11 scored. '
    'Grade schools have no enrolled grades 9–12 in the same-year October snapshot.')


def digest(path):
    return hashlib.sha256(Path(path).read_bytes()).hexdigest()


def code(value):
    if isinstance(value, (int, float)) and value == int(value):
        value = str(int(value))
    if not isinstance(value, str) or not re.fullmatch(r'\d{1,9}', value):
        raise ValueError(f'Invalid Pennsylvania authoritative ID: {value!r}')
    return value.zfill(9)


def school_id(aun, school):
    if code(school) == '000000000':
        raise ValueError('Zero building code is an intermediate unit aggregate, not a school')
    return code(aun) + '-' + code(school)


def numeric(value):
    if value in (None, '', '*', '**', '***', 'N/A', 'NA', 'Insufficient N', 'Insufficient Data'):
        return None
    value = float(value)
    if not math.isfinite(value):
        raise ValueError('Non-finite source number')
    return value


def count(value):
    value = numeric(value)
    if value is not None and (value < 0 or value != int(value)):
        raise ValueError('Invalid count')
    return None if value is None else int(value)


def grade_school(grades):
    # PKA/PKP/PKF, K4A/K4P/K4F, K5A/K5P/K5F, grades 1–12.
    if len(grades) != 21 or any(count(v) is None for v in grades):
        raise ValueError('Incomplete same-year enrolled grades')
    return not any(grades[17:])


def result(raw):
    if raw is None:
        return None, None, 'PSSA school total not reported'
    if len(raw) != 6:
        raise ValueError('Unexpected native PSSA school-total fields')
    tested = count(raw[0])
    rate = numeric(raw[5])
    if rate is None:
        return None, None, 'PDE suppressed native school-total proficiency (fewer than 11 scored)'
    levels = [numeric(v) for v in raw[1:5]]
    if tested is None or tested < 11 or any(v is None or not 0 <= v <= 100 for v in levels) or not 0 <= rate <= 100:
        raise ValueError('Published total lacks valid-score denominator or performance rates')
    if abs(sum(levels)-100) > .21 or abs(levels[0]+levels[1]-rate) > .11:
        raise ValueError('Native proficiency levels do not reconcile within source rounding')
    return rate, tested, None


def income_result(raw):
    if raw is None:
        return None, None, None, 'Same-year individual economic status not reported'
    if len(raw) != 3:
        raise ValueError('Unexpected low-income count fields')
    total, low = [count(v) for v in raw[:2]]
    rate = numeric(raw[2])
    if None in (total, low, rate):
        return total, None, None, 'Individual economic status suppressed or missing'
    # The workbook truncates the proportion to four decimals (e.g.71/996 -> .0712).
    if total <= 0 or low > total or not 0 <= rate <= 1 or not -.000001 <= low/total-rate <= .000101:
        raise ValueError('Published income percentage does not reconcile to count ratio')
    return total, low, 100*low/total, None


def worksheet(path, name):
    book = openpyxl.load_workbook(path, read_only=True, data_only=True)
    try:
        return list(book[name].values)
    finally:
        book.close()


def extract():
    sources = {key:dict(path='data/raw/'+filename,url=url,sha256=digest(ROOT/'data/raw'/filename)) for key,(filename,url) in FILES.items()}
    grades = worksheet(ROOT/sources['pa-grades-2025']['path'], 'LEA and School')
    income = worksheet(ROOT/sources['pa-income-2025']['path'], '2025 LIP by School')
    if grades[2][0] != 'Year: 2024-2025' or grades[4][15:28] != ('001','002','003','004','005','006','007','008','009','010','011','012','Total'):
        raise ValueError('Unexpected same-year enrollment schema')
    if income[3][0] != 'Year: 2024-2025' or income[4][5:8] != ('TOTAL ENROLLMENT','LOW-INCOME ENROLLMENT','PERCENT OF ENROLLMENT FROM LOW INCOME FAMILIES'):
        raise ValueError('Unexpected same-year individual income schema')
    profiles, excluded = {}, []
    for rownum,row in enumerate(grades[5:],6):
        if not isinstance(row[0],(int,float)):
            continue
        if code(row[4]) == '000000000':
            excluded.append(dict(source_row=rownum,raw=list(row),reason='Intermediate unit aggregate, not a building'))
            continue
        key = school_id(row[0],row[4])
        if key in profiles:
            raise ValueError('Duplicate school enrollment identity')
        values = list(row[6:27])
        if sum(values) != count(row[27]):
            raise ValueError('Grade counts disagree with total enrollment')
        profiles[key] = dict(school_id=key, district_id=code(row[0]),name=row[5],district=row[1],
            lea_type=row[2],county=row[3],identity_raw=list(row[:6]),grades=values,
            grade_school=grade_school(values),grade_source_row=rownum,income_raw=None,income_source_row=0)
    for rownum,row in enumerate(income[5:],6):
        if not isinstance(row[0],str) or not row[0].isdigit():
            continue
        if code(row[3]) == '000000000':
            continue
        key = school_id(row[0],row[3])
        if key not in profiles or profiles[key]['income_raw'] is not None:
            raise ValueError('Income lacks same-year school or duplicates ID')
        raw = list(row[5:8])
        total,_,_,_ = income_result(raw)
        if total != sum(profiles[key]['grades']):
            raise ValueError('Income and grade enrollments are not same-year matches')
        profiles[key].update(income_raw=raw,income_identity_raw=list(row[:5]),income_source_row=rownum)
    observations, unmatched = [], []
    book = openpyxl.load_workbook(ROOT/sources['pa-assessment-2025']['path'],read_only=True,data_only=True)
    sheet = book['PSSA']
    header = next(sheet.iter_rows(min_row=4,max_row=4,max_col=15,values_only=True))
    if header[6:10] != ('Subject','Group','Grade','Number Scored') or header[14] != 'Percent Proficient and above':
        raise ValueError('Unexpected native assessment schema')
    for rownum,row in enumerate(sheet.iter_rows(min_row=5,max_col=15,values_only=True),5):
        if row[7] != 'All Students' or row[8] != 'Total':
            continue
        if row[0] != YEAR or row[6] not in ['Math','English Language Arts']:
            raise ValueError('Wrong assessment scope/year')
        key = school_id(row[1],row[2])
        record = dict(school_id=key,subject='math' if row[6]=='Math' else 'reading',
            source_row=rownum,identity_raw=list(row[:9]),raw=list(row[9:15]))
        result(record['raw'])
        (observations if key in profiles else unmatched).append(record)
    book.close()
    EXTRACT.write_text(json.dumps(dict(state='PA',year=YEAR,school_year='2024-25',retrieved=str(date.today()),
        sources=sources,schools=list(profiles.values()),assessments=observations,
        excluded_enrollment=excluded,unmatched_assessments=unmatched),separators=(',',':'),allow_nan=False))


def import_data(db,payload):
    if (payload['state'],payload['year'],payload['school_year']) != ('PA',YEAR,'2024-25'):
        raise ValueError('Wrong source year/state')
    for table in ['model_run','assessment_observation','economic_observation','school','source']:
        db.execute(f'DELETE FROM {table} WHERE dataset_id=?',(DATASET,))
    db.execute('DELETE FROM assessment_definition WHERE id LIKE ?',(DATASET+':%',))
    db.execute('INSERT OR IGNORE INTO dataset VALUES (?,?,?,?,?)',(DATASET,'PA','Pennsylvania PSSA 2025','Native grade-school PSSA grades 3–8 totals','ready'))
    for key,source in dict(payload['sources'],**{'pa-extract':dict(path=str(EXTRACT.relative_to(ROOT)),url=PORTAL,sha256=digest(EXTRACT))}).items():
        db.execute('INSERT INTO source VALUES (?,?,?,?,?,?)',(key,DATASET,source['path'],source['url'],source['sha256'],payload['retrieved']))
    db.execute('INSERT OR REPLACE INTO economic_definition VALUES (?,?,?,?)',('pa-economic-2025','Pennsylvania individual reported economic disadvantage 2024–25',ECON_DEFINITION,FILES['pa-pims-2025'][1]))
    definition = definition_id(DATASET,YEAR,ASSESSMENT,'ES')
    db.execute('INSERT INTO assessment_definition VALUES (?,?,?,?,?,?,?,?)',(definition,'PA',ASSESSMENT,YEAR,'ES','3–8',ASSESSMENT_DEFINITION,FILES['pa-assessment-2025'][1]))
    profiles = {}
    for order,p in enumerate(payload['schools']):
        key = school_id(p['identity_raw'][0],p['identity_raw'][4])
        if key in profiles or key != p['school_id'] or p['district_id'] != code(p['identity_raw'][0]) or p['name'] != p['identity_raw'][5] or p['grade_school'] != grade_school(p['grades']):
            raise ValueError('Invalid public enrolled-school identity/scope')
        profiles[key]=p
        if not p['grade_school']:
            continue
        total,low,pct,_ = income_result(p['income_raw'])
        if p['income_raw'] is not None and (total != sum(p['grades']) or key != school_id(p['income_identity_raw'][0],p['income_identity_raw'][3])):
            raise ValueError('Income is not same-year school enrollment')
        db.execute('INSERT INTO school VALUES (?,?,?,?,?,?,?,?,?,?)',(DATASET,key,p['name'],p['district_id'],p['district'],None,p['county'],json.dumps(p),'pa-extract',order))
        db.execute('INSERT INTO economic_observation VALUES (?,?,?,?,?,?,?,?,?,?,?,?)',(DATASET,key,YEAR,'pa-economic-2025',p['name'],total,low,pct,'Individual reported economic disadvantage',json.dumps(p),'pa-income-2025',p['income_source_row']))
    observations={}
    for r in payload['assessments']:
        key=(r['school_id'],r['subject'])
        if key in observations or key[0] not in profiles or key[1] not in ['math','reading'] or key[0] != school_id(r['identity_raw'][1],r['identity_raw'][2]) or r['identity_raw'][0] != YEAR or r['identity_raw'][7:9] != ['All Students','Total']:
            raise ValueError('Duplicate or wrong native school-total identity')
        result(r['raw']);observations[key]=r
    for key,p in profiles.items():
        if not p['grade_school']:
            continue
        for subject in ['math','reading']:
            r=observations.get((key,subject),{})
            pct,tested,reason=result(r.get('raw'))
            db.execute('INSERT INTO assessment_observation VALUES (?,?,?,?,?,?,?,?,?,?,?)',(DATASET,key,definition,subject,pct,tested,
                'reported' if reason is None else 'suppressed_or_not_reported' if 'suppressed' in reason else 'not_reported',
                str(r['raw'][5]) if r else None,str(r['raw'][0]) if r else None,'pa-assessment-2025',r.get('source_row',0)))


def prepare(database=DEFAULT_DB):
    payload=json.loads(EXTRACT.read_text())
    with connect(database) as db:
        import_data(db,payload)
        assessments=pl.DataFrame([dict(r) for r in db.execute('''SELECT a.school_id,CAST(d.year AS TEXT) AS year,
            d.level,d.name AS assessment,a.subject,a.proficiency,a.tested FROM assessment_observation a
            JOIN assessment_definition d ON a.definition_id=d.id WHERE a.dataset_id=?''',(DATASET,))])
        incomes=pl.DataFrame([dict(r) for r in db.execute('''SELECT school_id,CAST(year AS TEXT) AS year,name,enrollment,
            low_income,percentage,source_label AS income_label FROM economic_observation WHERE dataset_id=?''',(DATASET,))])
        records,models=build_history(assessments,incomes);save_models(db,records,models,DATASET)
        if db.execute('PRAGMA foreign_key_check').fetchall():raise ValueError('Foreign key check failed')
    by_id={r['school_id']:r for r in records}
    lookup={(r['school_id'],r['subject']):r for r in payload['assessments']}
    schools=[]
    for p in payload['schools']:
        if not p['grade_school']:continue
        key=p['school_id'];annual=by_id[key];income_reason=income_result(p['income_raw'])[3]
        for subject in ['math','reading']:
            if subject not in annual['subjects']:
                assessment_reason=result(lookup.get((key,subject),{}).get('raw'))[2]
                annual['exclusions'][subject]='; '.join(filter(None,[income_reason,assessment_reason])) or annual['exclusions'].get(subject,'Model unavailable')
        if 'combined' not in annual['subjects']:
            annual['exclusions']['combined']='; '.join(f'{"ELA" if s=="reading" else "Math"}: {annual["exclusions"][s]}' for s in ['math','reading'] if s in annual['exclusions'])
        schools.append(dict(id=key,name=p['name'],short=p['name'],level='ES',program='Unclassified',programs=['Unclassified'],
            district=p['district'],city=None,county=p['county'],latitude=None,longitude=None,income=annual['income'],
            enrollment=annual['enrollment'],metrics=annual['subjects'],exclusions=annual['exclusions'],history=[annual]))
    methodology=dict(summary='Pennsylvania grade-school models use published PSSA grades 3–8 school totals and same-year individual reported economic status.',
        details=[ECON_DEFINITION,ASSESSMENT_DEFINITION,
            'Use the published All Students / Total proficiency percentage for each subject, including its source rounding. Number Scored is the valid-score denominator; no proficient student count is inferred from rounded percentages and no suppressed grade result is reconstructed.',
            'The October 1 grade-enrollment registry determines scope: schools enrolling any grades 9–12 are excluded. Income and grades match by nine-digit AUN plus school number in the same 2024–25 snapshot. Intermediate-unit aggregate code 000000000 is not a school.',
            'Externally studentized residuals use the complete eligible Pennsylvania grade-school population for each subject. Combined is the equally weighted mean of Math and ELA. Verified score denominators support sampling intervals. Associations do not estimate causal school effectiveness or overall quality.',
            'One year only. No authoritative school-ID coordinate crosswalk has been imported; schools are available in the list and charts while map locations remain unavailable.'],
        sources=[dict(label='PSSA native school totals',url=FILES['pa-assessment-2025'][1]),
                 dict(label='Same-year individual economic enrollment',url=FILES['pa-income-2025'][1]),
                 dict(label='Same-year enrolled grade counts',url=FILES['pa-grades-2025'][1]),
                 dict(label='PIMS 2024–25 manual, field 88',url=FILES['pa-pims-2025'][1])])
    output=dict(year='2024–25',assessment_year=YEAR,income_label='Individual reported economic disadvantage',
        levels=dict(ES=dict(year=YEAR,label='Grade schools · PSSA',assessment=ASSESSMENT,outcome='Proficient and above',math_label='Math',note='PSSA grades 3–8 · native school totals')),
        schools=schools,program_options=['Unclassified'],models=dict(ES={s:next((m for m in models if m['subject']==s),None) for s in ['math','reading','combined']}),
        history_years=[YEAR],history_models=models,methodology=methodology,
        coverage_note='Pennsylvania 2024–25 public grade schools: native All Students/Total PSSA grades 3–8, excluding late entrants, homeschoolers and source-excluded EL/LIFE students. PSSA only; PASA excluded. Same-year individual PIMS economic status, separate from CEP meal eligibility. No enrolled grades 9–12. School coordinates unavailable; one year only.')
    FOLDER.mkdir(parents=True,exist_ok=True)
    (FOLDER/'schools.json').write_text(json.dumps(output,separators=(',',':'),allow_nan=False))
    (FOLDER/'history.json').write_text(json.dumps(dict(records=records,models=models),separators=(',',':'),allow_nan=False))
    (FOLDER/'catalog.json').write_text(json.dumps(dict(schema_version=1,prepare_script='scripts/prepare_pennsylvania.py',state=dict(id='PA',name='Pennsylvania'),regions=[dict(id='pennsylvania',name='Grade schools',dataset=DATASET,status='ready',geography='pennsylvania',levels=['ES'],schools='data/pennsylvania/schools.json',boundaries=None,comparison='Pennsylvania · {year} · PSSA · {note}',model_scope='Public grade-school native PSSA grades 3–8 totals and same-year individual economic status',map_source='School coordinates unavailable; no authoritative coordinate crosswalk imported')]),indent=2)+'\n')
    audit=dict(year=YEAR,public_enrollment_schools=len(payload['schools']),grade_schools=len(schools),
        excluded_high_or_mixed=sum(not p['grade_school'] for p in payload['schools']),excluded_iu_aggregates=len(payload['excluded_enrollment']),
        unmatched_native_assessment_rows=len(payload['unmatched_assessments']),eligible={m['subject']:m['n'] for m in models},mapped=0,
        income_exclusions=dict(Counter(income_result(p['income_raw'])[3] for p in payload['schools'] if p['grade_school'] and income_result(p['income_raw'])[3])),
        assessment_exclusions={s:dict(Counter(result(lookup.get((p['school_id'],s),{}).get('raw'))[2] for p in payload['schools'] if p['grade_school'] and result(lookup.get((p['school_id'],s),{}).get('raw'))[2])) for s in ['math','reading']})
    (FOLDER/'coverage.json').write_text(json.dumps(audit,indent=2)+'\n');print(json.dumps(audit,indent=2))


if __name__=='__main__':
    parser=argparse.ArgumentParser(description=__doc__);parser.add_argument('--extract',action='store_true')
    parser.add_argument('--database',type=Path,default=DEFAULT_DB);args=parser.parse_args()
    if args.extract:extract()
    prepare(args.database)
