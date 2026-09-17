"""Convert official CPS workbooks into source CSVs; preserve suppression strings.
Usage: python scripts/import_assessments.py /path/iar.xlsx /path/sat.xlsx
Downloads and source URLs are documented in data/source/README.md.
"""
import csv
import sys
from pathlib import Path
import openpyxl
ROOT = Path(__file__).resolve().parents[1]

def convert(iar_path, sat_path):
    rows = []
    history_rows = []
    iar = openpyxl.load_workbook(iar_path, read_only=True, data_only=True)
    for subject, sheet in [('reading', 'IAR-PARCC ELA Results'), ('math', 'IAR-PARCC Math Results')]:
        for row in iar[sheet].iter_rows(min_row=3, values_only=True):
            if not isinstance(row[2], (int, float)) or not str(row[3]).startswith('Combined'):
                continue
            year = int(row[2])
            if row[3] == 'Combined ELA Grades 3-8' or row[3] == 'Combined Math Grades 3-8':
                level, assessment = 'ES', 'IAR/PARCC'
            elif row[3] == 'Combined ELA Grades 9-12' or row[3] == 'Combined Math Grades 9-12':
                level, assessment = 'HS', 'IAR/PARCC'
            else:
                continue
            history_rows.append([str(int(row[0])), year, level, assessment, subject, row[4], row[11]])
            if year == 2024 and level == 'ES':
                rows.append([str(int(row[0])), level, subject, row[4], row[11]])
    sat = openpyxl.load_workbook(sat_path, read_only=True, data_only=True)
    for row in sat['All Students Data'].iter_rows(min_row=2, values_only=True):
        if len(row) > 20 and row[0] == '2023-2024' and row[3] == 'SAT':
            for subject, n, pct in [('reading', 6, 15), ('math', 8, 20)]:
                rows.append([str(int(row[1])), 'HS', subject, row[n], row[pct]])
                history_rows.append([str(int(row[1])), 2024, 'HS', 'SAT', subject, row[n], row[pct]])
    with (ROOT / 'data/source/assessments-2024.csv').open('w', newline='') as f:
        writer = csv.writer(f)
        writer.writerow(['school_id', 'level', 'subject', 'tested', 'proficiency'])
        writer.writerows(rows)
    with (ROOT / 'data/source/assessments-history.csv').open('w', newline='') as f:
        writer = csv.writer(f)
        writer.writerow(['school_id', 'year', 'level', 'assessment', 'subject', 'tested', 'proficiency'])
        writer.writerows(history_rows)

if __name__ == '__main__':
    convert(*sys.argv[1:])
