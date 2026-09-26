"""NYC public aggregates -> canonical SQLite -> static JSON. No implicit downloads."""
import argparse
from collections import Counter, defaultdict
import hashlib
import json
import math
import re
import struct
import zipfile

import polars as pl
from database import ROOT, DEFAULT_DB, connect, add_source, definition_id, save_models
from prepare_data import build_history

DATASET = 'nyc'
EXTRACT = ROOT/'data/source/nyc.json'
BASE = 'https://infohub.nyced.org/docs/default-source/default-document-library/'
SOURCES = {
    'ela': ('nyc-school-ela.xlsx', BASE+'school-ela-results-public.xlsx'),
    'math': ('nyc-school-math.xlsx', BASE+'school-math-results-public.xlsx'),
    'income': ('nyc-demographics.xlsx', BASE+'demographic-snapshot-2021-22-to-2025-26-public.xlsx'),
    'regents': ('nyc-regents.xlsx', BASE+'2014-15-to-2022-23-nyc-regents-overall-and-by-category---public.xlsx'),
    'locations': ('nyc-locations.zip', 'https://data.cityofnewyork.us/download/jfju-ynrr/application/x-zip-compressed'),
    'boundary': ('nyc-boundary.geojson', 'https://data.cityofnewyork.us/resource/gthc-hcne.geojson'),
}
BOROUGHS = dict(M='Manhattan', X='Bronx', K='Brooklyn', Q='Queens', R='Staten Island')
HS_LEVELS = {'High school', 'Secondary School', 'K-12 all grades'}


def number(value):
    """Respect suppression/ranges, including bounded poverty values; never impute."""
    if value is None or str(value).strip().casefold() in {'', 's', 'na', 'n/a', '*', 'above 95%', 'below 5%'}:
        return None
    result = float(value)
    if not math.isfinite(result):
        raise ValueError(f'Non-finite value: {value}')
    return result


def validated_rate(raw):
    n, k, rate = [number(raw[key]) for key in ['tested', 'proficient', 'rate']]
    if n is not None and (n < 0 or n != int(n)):
        raise ValueError('Invalid tested count')
    if k is not None and (k < 0 or k != int(k) or n is None or k > n):
        raise ValueError('Invalid proficient count')
    if rate is not None and not 0 <= rate <= 100:
        raise ValueError('Invalid proficiency rate')
    if n and k is not None and rate is not None and abs(rate-100*k/n) > .011:
        raise ValueError('Rate does not reconcile with published numerator/denominator')
    # Never reconstruct a suppressed percentage from another field.
    return rate, int(n) if n is not None else None


def read_locations(path):
    """Read the official DBF's explicit latitude/longitude, joined by ATS/DBN."""
    with zipfile.ZipFile(path) as archive:
        members = [n for n in archive.namelist() if n.endswith('.dbf')]
        if len(members) != 1:
            raise ValueError('Expected one location table')
        b = archive.read(members[0])
    count, header, length = struct.unpack_from('<IHH', b, 4)
    fields = [(b[i:i+11].split(b'\0')[0].decode(), b[i+16]) for i in range(32, header-1, 32)]
    if not {'ATS', 'Latitude', 'Longitude'} <= {f for f, _ in fields}:
        raise ValueError('Location fields changed')
    if len(b) < header + count*length:
        raise ValueError('Truncated location table')
    candidates = defaultdict(set)
    for i in range(count):
        record = b[header+i*length:header+(i+1)*length]
        if record[:1] == b'*':
            continue
        offset, row = 1, {}
        for name, size in fields:
            row[name] = record[offset:offset+size].decode('utf-8').strip()
            offset += size
        lat, lon = number(row['Latitude']), number(row['Longitude'])
        if lat is not None and lon is not None and 40.4 <= lat <= 41 and -74.3 <= lon <= -73.65:
            candidates[row['ATS']].add((lat, lon))
    # Multiple campuses at different coordinates have no inferred single location.
    return {sid: dict(latitude=next(iter(points))[0], longitude=next(iter(points))[1])
            for sid, points in sorted(candidates.items()) if len(points) == 1}


def extract():
    import openpyxl
    source_meta, notes = {}, {}
    for key, (filename, url) in SOURCES.items():
        path = ROOT/'data/raw'/filename
        source_meta[key] = dict(path=str(path.relative_to(ROOT)), url=url, sha256=hashlib.sha256(path.read_bytes()).hexdigest())
    income = []
    w = openpyxl.load_workbook(ROOT/'data/raw'/SOURCES['income'][0], read_only=True, data_only=True)
    notes['income'] = [r[0] for r in w['NOTES'].values if r[0]]
    rows = w['School'].values; headers = next(rows)
    for row in rows:
        if not row[0]:
            continue
        r = dict(zip(headers, row)); sid = r['DBN']
        if not re.fullmatch(r'\d{2}[MKQRX]\d{3}', sid):
            raise ValueError(f'Unexpected DBN: {sid}')
        year = int(r['Year'].split('-')[0])+1
        income.append(dict(school_id=sid, year=year, name=r['School Name'], enrollment=r['Total Enrollment'],
            low_income=r['# Poverty'], fraction=r['% Poverty'],
            grades={str(g):r[f'Grade {g}'] for g in range(1,13)}, source='income'))
    w.close()
    keys = [(r['school_id'],r['year']) for r in income]
    if len(keys) != len(set(keys)):
        raise ValueError('Duplicate school-year demographics')
    years = {r['year'] for r in income}
    assessments = []
    for subject, key in [('reading','ela'), ('math','math')]:
        w = openpyxl.load_workbook(ROOT/'data/raw'/SOURCES[key][0], read_only=True, data_only=True)
        notes[key] = [r[0] for r in w['NOTES'].values if r[0]]
        rows = w['ELA - All' if key == 'ela' else 'Math - All'].values; headers = next(rows)
        for row in rows:
            r = dict(zip(headers,row))
            if r['Grade'] != 'All Grades' or r['Category'] != 'All Students' or r['Year'] not in years:
                continue
            year = int(r['Year'])
            assessments.append(dict(school_id=r['DBN'], name=r['School Name'], year=year, level='ES', subject=subject,
                assessment='NYSTP (2023 standards)' if year >= 2023 else 'NYSTP (2018 standards)',
                tested=r['Number Tested'], proficient=r['# Level 3+4'], rate=r['% Level 3+4'], source=key))
        w.close()
    w = openpyxl.load_workbook(ROOT/'data/raw'/SOURCES['regents'][0], read_only=True, data_only=True)
    notes['regents'] = [r[0] for r in w['Notes'].values if r[0]]
    rows = w['All Students'].values; headers = next(rows)
    for row in rows:
        r = dict(zip(headers,row))
        if r['Year'] not in years or r['Category'] != 'All Students' or r['School Level'] not in HS_LEVELS:
            continue
        if r['Regents Exam'] not in ['Common Core Algebra','Common Core English']:
            continue
        assessments.append(dict(school_id=r['School DBN'], name=r['School Name'], year=int(r['Year']), level='HS',
            subject='math' if r['Regents Exam']=='Common Core Algebra' else 'reading', assessment='Regents ELA / Algebra I',
            tested=r['Total Tested'], proficient=r['Number Scoring 65 or Above'], rate=r['Percent Scoring 65 or Above'], source='regents'))
    w.close()
    keys = [(r['school_id'],r['year'],r['level'],r['subject']) for r in assessments]
    if len(keys) != len(set(keys)):
        raise ValueError('Duplicate school/year/level/subject results')
    for r in assessments:
        validated_rate(r)
    payload = dict(retrieved='2026-09-26', sources=source_meta, notes=notes, income=income,
        assessments=sorted(assessments,key=lambda r:(r['school_id'],r['year'],r['level'],r['subject'])),
        locations=read_locations(ROOT/'data/raw'/SOURCES['locations'][0]))
    EXTRACT.write_text(json.dumps(payload,separators=(',',':'),allow_nan=False))
    boundary = json.loads((ROOT/'data/raw'/SOURCES['boundary'][0]).read_text())
    boundary.pop('crs',None)
    for f in boundary['features']:
        for polygon in f['geometry']['coordinates']:
            for i, ring in enumerate(polygon):
                area = sum(a[0]*b[1]-b[0]*a[1] for a,b in zip(ring,ring[1:]))
                if (i == 0 and area > 0) or (i > 0 and area < 0):
                    ring.reverse()
                ring[:] = [[round(x,5),round(y,5)] for x,y in ring]
    (ROOT/'data/nyc').mkdir(exist_ok=True)
    (ROOT/'data/nyc/boundary.geojson').write_text(json.dumps(boundary,separators=(',',':')))
    print(f'Extracted {len(income)} demographic and {len(assessments)} assessment records')


def import_data(db, payload):
    # Explicit deletions respect foreign keys and make each refresh reproducible.
    for table in ['model_run','assessment_observation','economic_observation','school','source']:
        db.execute(f'DELETE FROM {table} WHERE dataset_id=?',(DATASET,))
    db.execute('INSERT OR IGNORE INTO dataset VALUES (?,?,?,?,?)',
               (DATASET,'NY','New York City Public Schools','NYCPS published assessment cohorts','ready'))
    add_source(db,'nyc-extract',DATASET,EXTRACT,'https://infohub.nyced.org/reports/academics/test-results')
    for key, source in payload['sources'].items():
        db.execute('INSERT INTO source VALUES (?,?,?,?,?,?)',
                   (f'nyc-{key}',DATASET,source['path'],source['url'],source['sha256'],payload['retrieved']))
    profiles = {r['school_id']:r for r in sorted(payload['income'],key=lambda r:r['year'])}
    for r in payload['assessments']:
        profiles.setdefault(r['school_id'],dict(school_id=r['school_id'],name=r['name']))
    for order, (sid,p) in enumerate(sorted(profiles.items())):
        db.execute('INSERT INTO school VALUES (?,?,?,?,?,?,?,?,?,?)',
            (DATASET,sid,p['name'],sid[:2],f'NYC district {int(sid[:2])}',BOROUGHS[sid[2]],None,json.dumps(p),'nyc-extract',order))
    for order, r in enumerate(payload['income']):
        year=r['year']; eid=f'nyc-poverty-{year}'
        db.execute('INSERT OR IGNORE INTO economic_definition VALUES (?,?,?,?)',
            (eid,f'NYCPS Poverty {year}', 'NYCPS share qualifying for free/reduced-price meals or HRA benefits; bounded values stay unavailable. This is not the Economic Need Index.',payload['sources']['income']['url']))
        total, low, fraction = number(r['enrollment']),number(r['low_income']),number(r['fraction'])
        if fraction is not None and not 0 <= fraction <= 1:
            raise ValueError('Invalid poverty fraction')
        if total and low is not None and fraction is not None and abs(low/total-fraction) > .00001:
            raise ValueError('Poverty does not reconcile to enrollment')
        db.execute('INSERT INTO economic_observation VALUES (?,?,?,?,?,?,?,?,?,?,?,?)',
            (DATASET,r['school_id'],year,eid,r['name'],total,low,100*fraction if fraction is not None else None,'NYCPS Poverty',json.dumps(r),'nyc-income',order))
    for order, r in enumerate(payload['assessments']):
        did=definition_id(DATASET,r['year'],r['assessment'],r['level'])
        db.execute('INSERT OR IGNORE INTO assessment_definition VALUES (?,?,?,?,?,?,?,?)',
            (did,'NY',r['assessment'],r['year'],r['level'],'3–8' if r['level']=='ES' else 'High-school exam takers; multiple grades',
             'NYSTP Levels 3+4; 2023 standards break' if r['level']=='ES' else 'Regents ELA and Algebra I score 65+; highest score per student/exam/year; distinct subject populations',
             payload['sources'][r['source']]['url']))
        rate,tested=validated_rate(r)
        db.execute('INSERT INTO assessment_observation VALUES (?,?,?,?,?,?,?,?,?,?,?)',
            (DATASET,r['school_id'],did,r['subject'],rate,tested,'reported' if rate is not None else 'suppressed_or_not_reported',str(r['rate']),str(r['tested']),f"nyc-{r['source']}",order))


def prepare(database=DEFAULT_DB):
    payload=json.loads(EXTRACT.read_text())
    with connect(database) as db:
        import_data(db,payload)
        assessments=pl.DataFrame([dict(r) for r in db.execute("SELECT a.school_id,d.year,d.level,d.name assessment,a.subject,a.proficiency,a.tested FROM assessment_observation a JOIN assessment_definition d ON a.definition_id=d.id WHERE a.dataset_id='nyc'")],infer_schema_length=None)
        incomes=pl.DataFrame([dict(r) for r in db.execute("SELECT school_id,year,name,enrollment,low_income,percentage FROM economic_observation WHERE dataset_id='nyc'")],infer_schema_length=None)
        records,models=build_history(assessments,incomes)
        save_models(db,records,models,DATASET)
        profiles={r['school_id']:r for r in sorted(payload['income'],key=lambda r:r['year'])}
        snapshot={level:max(m['year'] for m in models if m['level']==level) for level in ['ES','HS']}
        histories=defaultdict(list)
        for r in records:
            histories[(r['school_id'],r['level'])].append(r)
        schools=[]
        for sid,p in sorted(profiles.items()):
            # Source assessment membership, not names, controls inclusion. Mixed-grade
            # schools are listed under high school; their ES results still fit ES models.
            level='HS' if any((number(p['grades'][str(g)]) or 0)>0 for g in range(9,13)) else 'ES'
            history=histories[(sid,level)]
            if not history:
                continue
            current=next((r for r in history if r['year']==snapshot[level]),None)
            current_income=next((r for r in payload['income'] if r['school_id']==sid and r['year']==snapshot[level]),None)
            fraction=number(current_income['fraction']) if current_income else None
            schools.append(dict(id=sid,name=p['name'],short=p['name'],level=level,program='Unclassified',
                district=f'NYC district {int(sid[:2])}',city=BOROUGHS[sid[2]],income=100*fraction if fraction is not None else None,
                enrollment=number(current_income['enrollment']) if current_income else None,
                latitude=payload['locations'].get(sid,{}).get('latitude'),longitude=payload['locations'].get(sid,{}).get('longitude'),
                metrics=current['subjects'] if current else {},history=history))
        level_meta={
            'ES':dict(year=snapshot['ES'],label='Grade schools · NYSTP',assessment='NYSTP',outcome='Proficiency',
                math_label='Math',note='NYSTP grades 3–8 · Levels 3 + 4'),
            'HS':dict(year=snapshot['HS'],label='High schools · Regents',assessment='Regents',outcome='Regents pass rate (65+)',
                math_label='Math · Algebra I',note='Regents ELA / Algebra I · score 65+ · exam takers, not a fixed grade cohort'),
        }
        exclusions=Counter(reason for r in records for reason in r['exclusions'].values())
        output=dict(year='2025–26 / 2022–23',assessment_year=snapshot['ES'],levels=level_meta,
            schools=schools,models={level:{m['subject']:m for m in models if m['year']==year and m['level']==level} for level,year in snapshot.items()},
            history_years=sorted({r['year'] for r in records}),history_models=models,
            coverage_note=f"NYCPS source coverage: {len(schools)} schools in the directory. Grade-school snapshot {snapshot['ES']}; Regents snapshot {snapshot['HS']}. Charter schools are excluded from the grade-school source; District 75 is generally excluded there. Only schools in the source cohorts are listed. Income is NYCPS Poverty, not Economic Need Index. Bounded poverty values are unavailable, never imputed. The demographic workbook omits schools closed before 2025–26, so historical models exclude schools without matching same-year income. Admissions classifications are unavailable. Locations use the August 2024 city school-point file; unmatched or ambiguous locations remain list-only.")
        folder=ROOT/'data/nyc';folder.mkdir(exist_ok=True)
        (folder/'schools.json').write_text(json.dumps(output,separators=(',',':'),allow_nan=False))
        (folder/'history.json').write_text(json.dumps(dict(records=records,models=models,exclusions=exclusions),separators=(',',':'),allow_nan=False))
        assert not db.execute('PRAGMA foreign_key_check').fetchall()
    print(json.dumps(dict(schools=len(schools),mapped=sum(s['latitude'] is not None for s in schools),snapshots=snapshot,
        models={level:{subject:m['n'] for subject,m in subjects.items()} for level,subjects in output['models'].items()},exclusions=exclusions),indent=2))


if __name__=='__main__':
    parser=argparse.ArgumentParser(description=__doc__)
    parser.add_argument('--extract',action='store_true')
    parser.add_argument('--database',default=DEFAULT_DB)
    args=parser.parse_args()
    if args.extract:
        extract()
    prepare(args.database)
