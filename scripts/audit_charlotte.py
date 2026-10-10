"""Audit Charlotte-Mecklenburg Schools' original 2024–25 source/cohort records without fitting."""
import argparse
from collections import Counter, defaultdict
import csv
import hashlib
import io
import json
from pathlib import Path
import re
import statistics
import subprocess
import zipfile

import openpyxl
from pypdf import PdfReader

import prepare_north_carolina as north_carolina
import plan_district_comparisons as planning
from state_registry import reported_count

ROOT = Path(__file__).resolve().parents[1]
OUTPUT = ROOT / 'data/source/charlotte-district-audit.json'
LEA = '3702970'
NATIVE_LEA = 'NC-600'
PSU = '600'
YEAR = 2025
ACADEMIC_YEAR = '2024-2025'
STATE_EXTRACT_SHA = '8d8cc1a75bbd58b1a73188aef9ad63c3008b9bf3d1334e4b3158e115caf70f28'
STATE_SNAPSHOT_SHA = 'fbaf4d121210b4a0b6f039d3693933f77f7c463168bb82dee880804cb46982c0'
RAW_INPUT_SHA = 'f9024f38c5a7bf179f94e0afa674fbb0ca8298b440c3352f70ff424abfd02c09'
SOURCES = {'dpi_performance': {'path': 'data/raw/nc-disaggregated-2025.zip', 'url': 'https://accrpt.tops.ncsu.edu/docs/disag_datasets/Disag_2024-25.zip', 'bytes': 21140473, 'sha256': '3db2e6a2e2a2d7c3ab41167750b16ddcf5a659f31fc3eee2a12b248857c65f0f'}, 'dpi_description': {'path': 'data/raw/nc-disaggregated-description-2025.pdf', 'url': 'https://accrpt.tops.ncsu.edu/docs/disag_datasets/Disag_2024-25.zip', 'bytes': 128241, 'sha256': '1c2a568a597dcd946872442637919d59c08e0f0b944b85d73317a5b713745bdf'}, 'dpi_income': {'path': 'data/raw/nc-eds-all-years.xlsx', 'url': 'https://www.dpi.nc.gov/eds-all-years-spreadsheet/open', 'bytes': 1328087, 'sha256': 'f34198b9ed0a8336e393e48ecd446c3d95d77469d72d1179174008d64e00a441'}, 'dpi_income_definition': {'path': 'data/raw/nc-eds-collection.pdf', 'url': 'https://www.dpi.nc.gov/documents/economically-disadvantaged-data-collection/open', 'bytes': 252019, 'sha256': 'c76951c64cc501a8ede2005f0111d490d563f6eeb011bbf7bb8fbf57f1ab5671'}, 'dpi_technical': {'path': 'data/raw/nc-school-grade-technical-2025.pdf', 'url': 'https://www.dpi.nc.gov/dpischoolgradetechnicalguidedraft2024-25/open', 'bytes': 1627575, 'sha256': '095a5d812dbd7a6f87c03dcc9f4eb69e943d569f22c803be7ab9ba52d5e7c049'}, 'ccd_directory': {'path': 'data/raw/ccd_sch_029_2425_w_1a_073025.zip', 'url': 'https://nces.ed.gov/ccd/Data/zip/ccd_sch_029_2425_w_1a_073025.zip', 'bytes': 13352819, 'sha256': '39326da788aa322353d20ceaf8ad4baed26272502cd05b066cf6c594988b21ab'}, 'ccd_membership': {'path': 'data/raw/ccd_sch_052_2425_l_1a_073025.zip', 'url': 'https://nces.ed.gov/ccd/Data/zip/ccd_sch_052_2425_l_1a_073025.zip', 'bytes': 212696691, 'sha256': '4a7f660c5fc5eaae488dd02fd43498f349fc828b227edd0970d5b6995ead4d4d'}, 'ccd_lea_directory': {'path': 'data/raw/ccd_lea_029_2425_w_1a_073025.zip', 'url': 'https://nces.ed.gov/ccd/Data/zip/ccd_lea_029_2425_w_1a_073025.zip', 'bytes': 2830266, 'sha256': '2745169e4bc7cd53adff179830c3d4b9bb035ac79580ecea01ef389a161936a9'}, 'ccd_lea_membership': {'path': 'data/raw/ccd_lea_052_2425_l_1a_073025.zip', 'url': 'https://nces.ed.gov/ccd/Data/zip/ccd_lea_052_2425_l_1a_073025.zip', 'bytes': 68497561, 'sha256': '501d72720a01c26e0e041cd3b1aa6653d0a94725ba0a1e85c42c4f183bc627ba'}, 'dpi_eds_web_definition': {'path': 'data/raw/nc-eds-definitions.html', 'url': 'https://www.dpi.nc.gov/data-reports/economically-disadvantaged', 'bytes': 63239, 'sha256': '36e318a6f5927bd14f53e714ac3527b92eab17abe6d0e889dbec0f0768c42df8'}}
SOURCES.update(dpi_src=dict(path='data/raw/nc-src-2025-2.zip',url='https://www.dpi.nc.gov/src-data-set-2024-252-2/open',bytes=69310425,sha256='1e7a1cbd3b189892dc49e53118b8b25213342a6aeca1658d7e44a6728c39a2d1'),
    dpi_src_dictionary=dict(path='data/raw/nc-src-dictionary-2025.xlsx',url='https://www.dpi.nc.gov/srcdatadictionary202425xlsx/open',bytes=73347,sha256='8865a7cb3fa74be71adc6f11be5615576a5ae84c15cd6a643af41ddcea5fe992'))
GRADES = ['PK', 'KG'] + [str(g) for g in range(1, 14)] + ['UG', 'AE']
LOWER_GRADES = ['Pre-Kindergarten', 'Kindergarten'] + ['Grade ' + str(g) for g in range(1, 9)]
OPERATIONAL = {'1', '3', '4', '5', '8'}
POLICY = {
    'native_location': 'Original same-year RCD location records and original dictionary retain exact native agency identities, raw category, school type, grade span and district aggregate. These supplemental classifications never replace CCD membership or native RG/GS assessment scope, infer provider/admissions roles, or add a cohort filter.',
    'roster': 'Complete original 2024–25 CCD LEA 3702970 / native NC-600 directory, including nonoperational/future schools, plus exact district membership total/grade records. Retain all native All Students assessment types, subjects and grades and every April PSU600 EDS row; district aggregates never become schools.',
    'identity': 'Cross-check same-year NCES LEA 3702970 and native NC-600; exact ST_SCHID NC-600-sss concatenates to six-character native school_code 600sss. NCES and native identifiers, not names or the first three characters alone, establish school identity.',
    'population': 'Prospective grade-school rows require complete unadjusted CCD PK–8 offerings, no high/adult grades, reported lower membership reconciliation, explicit literal-zero outside grades with Reported or Derived flags and native lower grade span. Exact Not Specified/Not Specified/Not Specified zero Derived rows establish no unallocated enrollment and retain their original flags; other Derived counts remain unverified. Preserve offered, enrolled, native spans, high, mixed, primary and unknown configurations separately. Discovery and fit quality never select members.',
    'assessment': 'Only native school_code individual school, subgroup ALL, type RG, grade GS and subject MA/RD published pct_glp may enter the prospective grade-school cohort. GS is grades 3–8 including grade8 NC Math1 EOC; GLP is Level3 and above. Native ALL/X1, EOC, subject EOG/ALL and grade rows remain provenance, never substitute or average.',
    'native_definition_limits': 'The description says tests completed in2024–25 include2024 summer school. It lists alternate type EXT1 while the actual native ZIP uses X1; both raw evidence and this codebook discrepancy remain explicit. Standalone Math1 excludes grade8 and is never substituted for native GS Math; the description note says9–12 while its subject/grade code definitions say9–13. These differences do not change the exact RG/GS selection.',
    'federal_denominator': 'The native source description says completed tests. The technical guide sections3.1–3.1.3 distinguish test-results reporting from School Performance Grades and long-term goals using the federal95% denominator adjustment. No SPG/long-term-goal rate enters this source cohort. num_tested valid-score scope remains independently unverified.',
    'eds_definition_limits': 'The original EDS webpage supports the individual direct/categorical eligibility examples, while collection PDF establishes optional local financial-need determinations and broad current-enrollment coverage. The webpage also contains stale CEP reimbursement-threshold discussion; that threshold is not used to define individual EDS or school income.',
    'income': north_carolina.INCOME_DEFINITION,
    'denominator': 'APR2025 den is the currently enrolled economic denominator. Fall CCD membership establishes grade scope only. Do not substitute fall enrollment, pct_nslp, adjusted reimbursement, tested EDS subgroup or inferred low-income counts.',
    'suppression': 'Exact published numeric percentage from 0 through100 is retained. Ranges, <5, >95, masked dash, stars and blanks remain unavailable. Do not reconstruct from other performance levels or income proxies.',
    'counts': 'num_tested remains raw unverified source evidence. Every valid-score denominator, sampling variance and interval remains null/unavailable; a ten-valid-scored minimum is uncertified. Published rates do not certify this count floor.',
    'combined': 'Equal mean of eligible native Math and ELA published GLP percentages; not the fraction proficient in both.',
    'approval': 'Source/cohort audit only: scope source_cohort_audit_only, status audit_pending and both source/modeling approvals false. No fits, canonical imports or browser releases. Independent numerical diagnostics, source decisions, integration and browser checks remain required; high/mixed/unknown populations remain unreleased.',
}
PLANNING = dict(source='data/source/district-comparison-candidates.json', operational_roster=186,
                enrollment=147299, potential_ES=146, potential_HS=25,
                explanation='Discovery only; exact original school identities and grade membership reconstruct the screen independently. Planning does not establish income/outcome usability or approve any district model.')


def fingerprint(value):
    return hashlib.sha256(json.dumps(value, sort_keys=True, separators=(',', ':'), ensure_ascii=False).encode()).hexdigest()


def load_north_carolina(path=north_carolina.EXTRACT):
    data = Path(path).read_bytes()
    if hashlib.sha256(data).hexdigest() != STATE_EXTRACT_SHA:
        raise ValueError('Complete North Carolina source fingerprint changed; fresh Charlotte-Mecklenburg audit required')
    payload = json.loads(data)
    north_carolina.validate(payload)
    return payload


def ccd_identity(raw, *, directory=False, district=False):
    if (raw.get('SCHOOL_YEAR') != ACADEMIC_YEAR or raw.get('ST') != 'NC'
            or raw.get('LEAID') != LEA or raw.get('ST_LEAID') != NATIVE_LEA):
        raise ValueError('Charlotte-Mecklenburg requires exact same-year NCES/native district identity')
    if district:
        return LEA
    if (not re.fullmatch(LEA + r'\d{5}', raw.get('NCESSCH', ''))
            or not re.fullmatch(r'NC-600-\d{3}', raw.get('ST_SCHID', ''))):
        raise ValueError('Charlotte-Mecklenburg requires exact authoritative school identity')
    if directory and raw.get('UPDATED_STATUS') not in {str(i) for i in range(1, 9)}:
        raise ValueError('Unrecognized same-year CCD school status')
    return PSU + raw['ST_SCHID'][-3:]


def native_identity(raw, *, income=False):
    key = raw.get('school_code', '')
    if not re.fullmatch(r'600(?:\d{3}|LEA)', key):
        raise ValueError('Foreign or malformed native Charlotte-Mecklenburg school/aggregate identity')
    if income:
        if (raw.get('psu_code') != PSU or raw.get('reporting_year') != '2025'
                or raw.get('collection_code') != 'APR'):
            raise ValueError('Charlotte-Mecklenburg EDS must be exact APR2025 PSU600')
    elif raw.get('subgroup') != 'ALL':
        raise ValueError('Charlotte-Mecklenburg raw assessment inventory must be native All Students')
    return key


def location_identity(raw):
    key=raw.get('agency_code')
    if type(raw.get('year'))is not int or raw['year']!=YEAR or raw.get('state')!='NC' or not isinstance(key,str)or not re.fullmatch(PSU+r'(?:\d{3}|LEA)',key):
        raise ValueError('Original RCD identity must be same-year exact native district')
    if key==PSU+'LEA':
        if raw.get('agency_level')!='LEA' or raw.get('lea_code')!='':
            raise ValueError('Original RCD aggregate identity/level is invalid')
    elif raw.get('agency_level')!='SCH' or raw.get('lea_code')!=PSU+'LEA':
        raise ValueError('Original RCD school identity/ownership is invalid')
    return key


def offered_evidence(raw):
    flags = {g: raw.get('G_' + g + '_OFFERED') for g in GRADES}
    complete = raw.get('IGOFFERED') == 'As reported' and raw.get('NOGRADES') == 'No' and all(v in {'Yes', 'No'} for v in flags.values())
    offered = [g for g, v in flags.items() if v == 'Yes']
    lower = any(str(g) in offered for g in range(3, 9))
    high = any(str(g) in offered for g in range(9, 14))
    band = ('unknown' if not complete else 'other' if any(g in offered for g in ['AE', 'UG'])
            else 'mixed' if lower and high else 'grade' if lower else 'high' if high else 'primary')
    return dict(flags=flags, complete_reported_flags=complete, offered=offered, band=band,
                offered_tested_grades=[str(g) for g in range(3, 9) if str(g) in offered],
                outside_PK12_offered=[g for g in ['13', 'AE', 'UG'] if g in offered],
                low=raw.get('GSLO'), high=raw.get('GSHI'), level=raw.get('LEVEL'))


def grade_school_contract(raw, total, grades):
    # Recompute the source grade contract from full original membership records;
    # ordinary offline replay never touches a membership cache or a database.
    if (raw.get('IGOFFERED') != 'As reported' or raw.get('NOGRADES') != 'No'
            or not re.fullmatch(r'0[1-8]', raw.get('GSHI', ''))
            or any(raw.get('G_' + g + '_OFFERED') not in {'Yes', 'No'} for g in GRADES)
            or any(raw.get('G_' + g + '_OFFERED') == 'Yes' for g in [str(i) for i in range(9, 14)] + ['AE'])
            or not grades):
        return False
    values = [reported_count(r) for g, r in grades.items() if g in LOWER_GRADES]
    if not values or any(v is None for v in values) or sum(values) != reported_count(total):
        return False
    if any(r['STUDENT_COUNT'] != '0' or r['DMS_FLAG'] not in {'Reported', 'Derived'} for g, r in grades.items() if g not in LOWER_GRADES):
        return False
    return raw['G_UG_OFFERED'] != 'Yes' or reported_count(grades.get('Ungraded')) == 0


def enrolled_evidence(raw, total_entry, grade_entries):
    total = None if total_entry is None else total_entry['raw']
    grades = {g: e['raw'] for g, e in grade_entries.items()}
    rows = [dict(grade=g, source_row=e['source_row'], raw_count=e['raw']['STUDENT_COUNT'],
                 flag=e['raw']['DMS_FLAG'], reported_count=grade_count(e['raw'])) for g, e in sorted(grade_entries.items())]
    lower = [reported_count(r) for g, r in grades.items() if g in LOWER_GRADES]
    known = LOWER_GRADES + ['Grade ' + str(g) for g in range(9, 14)] + ['Adult Education', 'Ungraded', 'Not Specified']
    positive = [r['grade'] for r in rows if r['reported_count'] is not None and r['reported_count'] > 0]
    native = grade_school_contract(raw, total, grades)
    if native != north_carolina.grade_school(raw, grades, total):
        raise ValueError('Independent grade contract disagrees with audited statewide contract')
    return dict(native_grade_school_contract=native, reported_total=reported_count(total),
                total_source_row=None if total_entry is None else total_entry['source_row'],
                total_raw=None if total_entry is None else total_entry['raw'], grade_values=rows,
                positive_enrolled_grades=positive, positive_unrecognized_grades=[g for g in positive if g not in known],
                lower_sum=sum(lower) if lower and all(v is not None for v in lower) else None,
                full_grade_sum=sum(r['reported_count'] for r in rows) if rows and all(r['reported_count'] is not None for r in rows) else None,
                missing_or_suppressed_grade_records=[r['grade'] for r in rows if r['reported_count'] is None])


def grade_count(raw):
    # CCD's exact zero Derived Not Specified row establishes no unallocated
    # enrollment; it is not a suppressed grade or a verified tested count.
    if (raw.get('GRADE') == 'Not Specified' and raw.get('STUDENT_COUNT') == '0'
            and raw.get('DMS_FLAG') == 'Derived' and raw.get('RACE_ETHNICITY') == 'Not Specified'
            and raw.get('SEX') == 'Not Specified'):
        return 0
    return reported_count(raw)


def income_evidence(entry):
    raw = None if entry is None else entry['raw']
    if raw is not None:
        native_identity(raw, income=True)
    pct = None if raw is None else north_carolina.exact_percent(raw['pct_eds'])
    den = None if raw is None else north_carolina.exact_enrollment(raw['den'])
    reason = ('Missing exact APR2025 individual school EDS row' if raw is None else
              'April individual EDS percentage is missing, ranged or suppressed' if pct is None else
              'April economic enrollment is missing, ranged, suppressed or zero' if den is None else None)
    return dict(source='dpi_income', source_row=None if entry is None else entry['source_row'],
                collection='APR 2025', percentage=pct, april_enrollment=den, verified_low_income_count=None,
                raw_percentage=None if raw is None else raw['pct_eds'], raw_enrollment=None if raw is None else raw['den'],
                exclusion=reason, definition=POLICY['income'])


def subject_evidence(entries, *, operational, configuration, offered, income):
    values = {}
    for subject in ['math', 'reading']:
        entry = entries.get(subject); raw = None if entry is None else entry['raw']
        if raw is not None and (raw['type'], raw['grade'], raw['subgroup'], raw['subject']) != ('RG', 'GS', 'ALL', 'MA' if subject == 'math' else 'RD'):
            raise ValueError('Prospective Charlotte-Mecklenburg outcome is not exact RG/GS/ALL Math/ELA')
        rate = None if raw is None else north_carolina.exact_percent(raw['pct_glp'])
        span_ok = raw is not None and north_carolina.native_grade_school(raw)
        applicable = operational and configuration and bool(offered['offered_tested_grades']) and span_ok
        scope_reason = ('Same-year CCD school is nonoperational' if not operational else
                        'CCD configuration is high, mixed, primary, adjusted or unknown' if not configuration or not offered['offered_tested_grades'] else
                        'Native lower-grade span is unavailable or conflicts with grade-school scope' if raw is not None and not span_ok else None)
        outcome_reason = ('Missing exact native regular GS All Students subject row' if raw is None else
                          'Native published GLP is missing, ranged or suppressed' if rate is None else None)
        reason = scope_reason or income['exclusion'] or outcome_reason
        values[subject] = dict(source='dpi_performance', source_row=None if entry is None else entry['source_row'],
            applicable=applicable, usable=reason is None, native_rate=rate,
            raw_rate=None if raw is None else raw['pct_glp'], raw_tested=None if raw is None else raw['num_tested'],
            native_grade_span=None if raw is None else raw['grade_span'], native_grade_span_compatible=span_ok,
            scope_exclusion=scope_reason, outcome_exclusion=outcome_reason, income_exclusion=income['exclusion'], exclusion=reason,
            valid_scores=None, sampling_variance=None, sampling_interval_available=False, ten_valid_scored_floor_certified=False)
    usable = all(v['usable'] for v in values.values())
    values['combined'] = dict(applicable=all(v['applicable'] for v in values.values()), usable=usable,
        native_rate=(values['math']['native_rate'] + values['reading']['native_rate']) / 2 if usable else None,
        exclusion='; '.join(s + ': ' + v['exclusion'] for s, v in values.items() if v['exclusion']) or None,
        valid_scores=None, sampling_variance=None, sampling_interval_available=False,
        ten_valid_scored_floor_certified=False, definition=POLICY['combined'])
    return values


def summarize(links, predicate, basis):
    profiles = [r for r in links if predicate(r)]
    subjects = {}
    for subject in ['math', 'reading', 'combined']:
        eligible = [r for r in profiles if r['subjects'][subject]['usable']]
        incomes = [r['income']['percentage'] for r in eligible]
        subjects[subject] = dict(applicable=sum(r['subjects'][subject]['applicable'] for r in profiles), usable=len(eligible),
            eligible_school_ids=[r['school_id'] for r in eligible], eligible_ncessch=[r['ncessch'] for r in eligible],
            prospective_30_school_floor_met=len(eligible) >= 30,
            exclusions=dict(sorted(Counter(r['subjects'][subject]['exclusion'] for r in profiles if not r['subjects'][subject]['usable']).items())),
            verified_valid_score_counts=0, ten_valid_scored_floor_certified=False, model_sampling_intervals_available=False,
            income=dict(minimum=min(incomes) if incomes else None, maximum=max(incomes) if incomes else None,
                        distinct=len(set(incomes)), sample_standard_deviation=statistics.stdev(incomes) if len(incomes) > 1 else None))
    return dict(profiles=len(profiles), school_ids=[r['school_id'] for r in profiles], membership_basis=basis, subjects=subjects)


def derive(document, snapshot):
    if fingerprint(snapshot) != STATE_SNAPSHOT_SHA:
        raise ValueError('Complete pinned North Carolina snapshot changed')
    if (type(document.get('schema_version')) is not int or document.get('schema_version') != 1
            or document.get('state') != 'NC' or type(document.get('year')) is not int or document.get('year') != YEAR
            or document.get('academic_year') != ACADEMIC_YEAR or document.get('nces_lea_id') != LEA or document.get('native_lea_id') != NATIVE_LEA
            or document.get('status') != 'audit_pending' or document.get('scope') != 'source_cohort_audit_only'
            or document.get('approved_for_source') is not False or document.get('approved_for_modeling') is not False
            or fingerprint(document.get('sources')) != fingerprint(SOURCES) or fingerprint(document.get('policy')) != fingerprint(POLICY)
            or fingerprint(document.get('planning')) != fingerprint(PLANNING)
            or document.get('state_extract') != dict(path='data/source/north-carolina.json', sha256=STATE_EXTRACT_SHA)):
        raise ValueError('Charlotte-Mecklenburg source-only identity/provenance/policy changed')
    inputs = document['raw_inputs']
    if fingerprint(inputs) != RAW_INPUT_SHA:
        raise ValueError('Complete retained original Charlotte-Mecklenburg records, definitions, headers or rows changed')
    for name in ['directory', 'membership', 'lea_directory', 'lea_membership', 'assessment', 'income', 'native_location']:
        section = inputs[name]
        if (type(section['header_source_row']) is not int or section['header_source_row'] != 1
                or len(section['header']) != len(set(section['header']))):
            raise ValueError('Original source header reference is invalid')
        seen = set()
        for e in section['records']:
            if type(e['source_row']) is not int or e['source_row'] < 2 or e['source_row'] in seen or set(e['raw']) != set(section['header']):
                raise ValueError('Original source row/field coverage is invalid')
            seen.add(e['source_row'])
    directory, ncids = {}, {}
    for entry in inputs['directory']['records']:
        raw = entry['raw']; key = ccd_identity(raw, directory=True)
        if key in directory or raw['NCESSCH'] in ncids:
            raise ValueError('Duplicate exact Charlotte-Mecklenburg directory identity')
        directory[key] = entry; ncids[raw['NCESSCH']] = key
    location={}
    for entry in inputs['native_location']['records']:
        key=location_identity(entry['raw'])
        if key in location:
            raise ValueError('Duplicate same-year original RCD location identity')
        location[key]=entry
    if set(location)!=(set(directory)|{PSU+'LEA'}):
        raise ValueError('Complete original RCD and CCD identities do not reconcile')
    totals, grades = {}, defaultdict(dict)
    for entry in inputs['membership']['records']:
        raw = entry['raw']; key = ccd_identity(raw)
        if key not in directory or directory[key]['raw']['NCESSCH'] != raw['NCESSCH']:
            raise ValueError('Membership row is outside exact school directory')
        reported_count(raw)
        if raw['TOTAL_INDICATOR'] == 'Education Unit Total':
            if key in totals or any(raw[k] != 'No Category Codes' for k in ['GRADE', 'RACE_ETHNICITY', 'SEX']):
                raise ValueError('Duplicate or non-total school membership population')
            totals[key] = entry
        elif raw['TOTAL_INDICATOR'] == 'Subtotal 4 - By Grade':
            unspecified_zero = (raw['GRADE'] == 'Not Specified' and raw['STUDENT_COUNT'] == '0' and raw['DMS_FLAG'] == 'Derived'
                                and raw['RACE_ETHNICITY'] == 'Not Specified' and raw['SEX'] == 'Not Specified')
            if raw['GRADE'] in grades[key] or not unspecified_zero and any(raw[k] != 'No Category Codes' for k in ['RACE_ETHNICITY', 'SEX']):
                raise ValueError('Duplicate or demographic grade membership population')
            grades[key][raw['GRADE']] = entry
        else:
            raise ValueError('Unexpected native school membership population')
    if len(inputs['lea_directory']['records']) != 1:
        raise ValueError('Retain exactly one original LEA directory row')
    ccd_identity(inputs['lea_directory']['records'][0]['raw'],district=True)
    lea_totals,lea_grades=[],{}
    for entry in inputs['lea_membership']['records']:
        raw=entry['raw'];ccd_identity(raw,district=True)
        if raw['TOTAL_INDICATOR']=='Education Unit Total':
            if any(raw[k]!='No Category Codes' for k in ['GRADE','RACE_ETHNICITY','SEX']):
                raise ValueError('Original LEA total has a foreign demographic population')
            lea_totals.append(entry)
        elif raw['TOTAL_INDICATOR']=='Subtotal 4 - By Grade':
            unspecified=(raw['GRADE']=='Not Specified'and raw['STUDENT_COUNT']=='0'and raw['DMS_FLAG']=='Derived'and raw['RACE_ETHNICITY']=='Not Specified'and raw['SEX']=='Not Specified')
            if raw['GRADE']in lea_grades or not unspecified and any(raw[k]!='No Category Codes' for k in ['RACE_ETHNICITY','SEX']):
                raise ValueError('Duplicate or demographic LEA grade population')
            lea_grades[raw['GRADE']]=entry
        else:
            raise ValueError('Unexpected original LEA membership population')
    if len(lea_totals)!=1 or not lea_grades:
        raise ValueError('Original LEA membership requires one total and complete grade rows')
    lea_total=lea_totals[0]
    if reported_count(lea_total['raw']) != PLANNING['enrollment']:
        raise ValueError('Original district total differs from discovery screen')
    lea_values=[grade_count(e['raw'])for e in lea_grades.values()]
    if any(v is None for v in lea_values)or sum(lea_values)!=reported_count(lea_total['raw']):
        raise ValueError('Original LEA grade rows do not reconcile to reported total')
    income, native, native_inventory = {}, defaultdict(dict), defaultdict(list)
    for entry in inputs['income']['records']:
        key = native_identity(entry['raw'], income=True)
        if key in income:
            raise ValueError('Duplicate April individual EDS school/aggregate')
        income[key] = entry
    unique_native = set()
    for entry in inputs['assessment']['records']:
        raw = entry['raw']; key = native_identity(raw)
        identity = (key, raw['subject'], raw['grade'], raw['type'], raw['subgroup'])
        if identity in unique_native:
            raise ValueError('Duplicate native assessment school/population row')
        unique_native.add(identity); native_inventory[key].append(entry)
        if raw['type'] == 'RG' and raw['grade'] == 'GS' and raw['subject'] in ['MA', 'RD']:
            native[key]['math' if raw['subject'] == 'MA' else 'reading'] = entry
    profiles = {s['id']: s for s in snapshot['schools'] if s['district_id'] == PSU}
    links, nonoperational = [], []
    for key, entry in sorted(directory.items()):
        raw = entry['raw']; op = raw['UPDATED_STATUS'] in OPERATIONAL
        offered = offered_evidence(raw); enrolled = enrolled_evidence(raw, totals.get(key), grades[key])
        eds = income_evidence(income.get(key))
        subjects = subject_evidence(native[key], operational=op, configuration=enrolled['native_grade_school_contract'], offered=offered, income=eds)
        pair = set(native[key]) == {'math', 'reading'}
        source_profile = (op and enrolled['native_grade_school_contract'] and pair and eds['exclusion'] is None
                          and all(v['native_grade_span_compatible'] for s, v in subjects.items() if s != 'combined'))
        link = dict(school_id=key, ncessch=raw['NCESSCH'], native_school_id=raw['ST_SCHID'], name=raw['SCH_NAME'],
            directory_source_row=entry['source_row'], operational=op, offered=offered, enrolled=enrolled,
            native_location=location[key],
            ccd_flags=dict(status=raw['UPDATED_STATUS'], status_text=raw['UPDATED_STATUS_TEXT'], school_type=raw['SCH_TYPE_TEXT'], charter=raw['CHARTER_TEXT'], virtual_status=None),
            income=eds, subjects=subjects, native_regular_gs_pair=pair,
            native_regular_gs_rows={s: e['source_row'] for s, e in sorted(native[key].items())},
            native_all_students_rows=[e['source_row'] for e in native_inventory.get(key, [])],
            native_grade_spans=sorted({e['raw']['grade_span'] for e in native_inventory.get(key, [])}),
            april_minus_fall_enrollment=None if eds['april_enrollment'] is None or enrolled['reported_total'] is None else eds['april_enrollment'] - enrolled['reported_total'],
            historical_state_source_profile=key in profiles, source_profile_eligible=source_profile)
        if source_profile != (key in profiles):
            raise ValueError('Original prospective source membership differs from complete pinned NC extract')
        if source_profile:
            old = profiles[key]
            if fingerprint([old['income'], old['enrollment'], old['low_income'], old['income_raw'], old['income_row'], old['directory_raw']]) != fingerprint([eds['percentage'], eds['april_enrollment'], None, income[key]['raw'], income[key]['source_row'], raw]):
                raise ValueError('Original April income/directory no longer matches pinned state subset')
            for subject, outcome in old['outcomes'].items():
                if fingerprint([outcome['raw'], outcome['source_row'], outcome['proficiency'], outcome['tested']]) != fingerprint([native[key][subject]['raw'], native[key][subject]['source_row'], subjects[subject]['native_rate'], None]):
                    raise ValueError('Original native RG/GS outcome differs from pinned state subset')
        (links if op else nonoperational).append(link)
    potential = dict(ES=[], HS=[])
    for link in links:
        key = link['school_id']; level = planning.potential_cohort(directory[key]['raw'], None if key not in totals else totals[key]['raw'], {g: e['raw'] for g, e in grades[key].items()})
        if level:
            potential[level].append(link['ncessch'])
    potential = {k: sorted(v) for k, v in potential.items()}
    recorded = inputs['planning']['record']
    if (recorded['nces_lea_id'] != LEA or recorded['native_lea_id'] != NATIVE_LEA or recorded['state'] != 'NC'
            or recorded['school_directory_operational_count'] != len(links)
            or recorded['enrollment'] != reported_count(lea_total['raw'])
            or recorded['potential_school_ids'] != potential or recorded['potential_cohorts'] != {k: len(v) for k, v in potential.items()}
            or recorded['membership_raw'] != lea_total['raw'] or recorded['membership_source_row'] != lea_total['source_row']
            or recorded['directory_source_row'] != inputs['lea_directory']['records'][0]['source_row']
            or any(inputs['lea_directory']['records'][0]['raw'][k] != v for k, v in recorded['directory_raw'].items())):
        raise ValueError('Original exact LEA/school records do not reconstruct discovery screen')
    cohorts = {
        'ccd_lower_membership_contracts': summarize(links,lambda r:r['enrolled']['native_grade_school_contract'],'Complete reconciled lower membership contract; may include primary-only schools before tested-grade offers are checked.'),
        'ccd_grade_school_configurations': summarize(links, lambda r: r['enrolled']['native_grade_school_contract'] and bool(r['offered']['offered_tested_grades']), 'Complete CCD lower membership plus at least one explicitly offered tested grade3–8, before income/outcome availability.'),
        'native_grade_school_source_profiles': summarize(links, lambda r: r['source_profile_eligible'], 'Independent exact original native RG/GS/ALL pair, compatible native/CCD grade scope and usable APR2025 individual EDS.'),
        'high_school_configurations': summarize(links, lambda r: r['offered']['band'] == 'high', 'Descriptive exact high offerings only; native EOC assessment population remains unaudited and no high model is approved.'),
        'mixed_configurations': summarize(links, lambda r: r['offered']['band'] == 'mixed', 'Descriptive mixed offerings; native GS rows never approve a mixed model.'),
        'primary_configurations': summarize(links, lambda r: r['offered']['band'] == 'primary', 'Descriptive primary-only offerings with no grade3–8 tested scope.'),
        'unknown_or_other_configurations': summarize(links, lambda r: r['offered']['band'] in {'unknown', 'other'}, 'Uncertain or outside grade offerings remain separate and unreleased.'),
    }
    coverage = dict(directory_records=len(directory), operational_schools=len(links), nonoperational_schools=len(nonoperational),
        operational_statuses=dict(sorted(Counter(r['ccd_flags']['status_text'] for r in links).items())),
        nonoperational_statuses=dict(sorted(Counter(r['ccd_flags']['status_text'] for r in nonoperational).items())),
        operational_offered_bands=dict(sorted(Counter(r['offered']['band'] for r in links).items())),
        operational_grade13_high_schools=[r['school_id'] for r in links if r['offered']['band'] == 'high' and '13' in r['offered']['offered']],
        operational_grade13_mixed_schools=[r['school_id'] for r in links if r['offered']['band'] == 'mixed' and '13' in r['offered']['offered']],
        operational_school_types=dict(sorted(Counter(r['ccd_flags']['school_type'] for r in links).items())),
        operational_charter=dict(sorted(Counter(r['ccd_flags']['charter'] for r in links).items())),
        school_membership_totals=len(totals), school_membership_grade_maps=len([g for g in grades.values() if g]),
        lea_fall_enrollment=reported_count(lea_total['raw']),
        lea_grade_membership_evidence=dict(total_source_row=lea_total['source_row'],total_raw=lea_total['raw'],grade_rows=[dict(grade=g,source_row=e['source_row'],raw=e['raw'],reported_count=grade_count(e['raw']))for g,e in sorted(lea_grades.items())],full_grade_sum=sum(lea_values)),
        operational_school_fall_enrollment=sum(r['enrolled']['reported_total'] for r in links if r['enrolled']['reported_total'] is not None),
        lea_minus_operational_school_fall_enrollment=reported_count(lea_total['raw']) - sum(r['enrolled']['reported_total'] for r in links if r['enrolled']['reported_total'] is not None),
        original_native_all_students_rows=len(inputs['assessment']['records']), original_native_units=len(native_inventory),
        original_native_location_records=len(location),original_native_location_school_records=sum(e['raw']['agency_level']=='SCH'for e in location.values()),
        native_location_categories=dict(sorted(Counter(e['raw']['category_code']for e in location.values()if e['raw']['agency_level']=='SCH').items())),
        native_location_school_types=dict(sorted(Counter(e['raw']['school_type']for e in location.values()if e['raw']['agency_level']=='SCH').items())),
        ccd_lower_membership_contracts=cohorts['ccd_lower_membership_contracts']['profiles'],
        native_assessment_absent_school_ids=[r['school_id']for r in links if not r['native_all_students_rows']],
        native_assessment_types=dict(sorted(Counter(e['raw']['type'] for e in inputs['assessment']['records']).items())),
        original_april_eds_rows=len(income), operational_usable_april_income=sum(r['income']['exclusion'] is None for r in links),
        april_district_aggregate_enrollment=income_evidence(income.get('600LEA'))['april_enrollment'],
        operational_school_april_enrollment=sum(r['income']['april_enrollment'] for r in links if r['income']['april_enrollment'] is not None),
        historical_state_source_profiles=len(profiles), prospective_source_profiles=cohorts['native_grade_school_source_profiles']['profiles'],
        grade_school_income_exclusions=dict(sorted(Counter(r['income']['exclusion'] for r in links if r['enrolled']['native_grade_school_contract'] and r['income']['exclusion']).items())),
        april_enrollment_different_from_fall=sum(r['april_minus_fall_enrollment'] not in {None, 0} for r in links),
        verified_valid_score_counts=0, sampling_intervals_available=False, ten_valid_scored_floor_certified=False)
    return dict(roster_links=links, nonoperational_directory=nonoperational, cohorts=cohorts, coverage=coverage,
        outside_directory_records=({name: [e for e in inputs[name]['records'] if e['raw']['school_code'] not in directory] for name in ['assessment', 'income']}|{'native_location':[e for e in inputs['native_location']['records']if e['raw']['agency_code']not in directory]}),
        source_definition_evidence=inputs['definitions'],
        planning_identity_reconciliation=dict(potential_school_ids=potential,
            source_eligible_ncessch=[r['ncessch'] for r in links if r['source_profile_eligible']],
            planning_ES_not_source_eligible=sorted(set(potential['ES']) - {r['ncessch'] for r in links if r['source_profile_eligible']}),
            source_not_planning_ES=sorted({r['ncessch'] for r in links if r['source_profile_eligible']} - set(potential['ES'])),
            combined_eligible_ncessch=cohorts['native_grade_school_source_profiles']['subjects']['combined']['eligible_ncessch'],
            planning_ES_not_combined_eligible=sorted(set(potential['ES']) - set(cohorts['native_grade_school_source_profiles']['subjects']['combined']['eligible_ncessch'])),
            high_grade13_not_planning_HS=[r['ncessch'] for r in links if r['offered']['band'] == 'high' and r['ncessch'] not in potential['HS']],
            explanation='Independent original size-screen reconstruction is descriptive only. Income and native grade/assessment scope establish source profiles separately; subject masks establish separate prospective outcomes. High, mixed, primary and uncertain configurations remain preserved outside the grade-school scope and do not approve a model.'),
        historical_state_source_approval=dict(approved_point_only=snapshot['source_audit']['approved_point_only'], district_source_approval=False, district_modeling_approval=False))


def validate(document, snapshot=None):
    expected = derive(document, snapshot if snapshot is not None else load_north_carolina())
    base = {'schema_version', 'state', 'year', 'academic_year', 'nces_lea_id', 'native_lea_id', 'status', 'scope',
            'approved_for_source', 'approved_for_modeling', 'sources', 'policy', 'planning', 'raw_inputs', 'state_extract'}
    if set(document) != base | set(expected) or any(fingerprint(document.get(k)) != fingerprint(v) for k, v in expected.items()):
        raise ValueError('Charlotte-Mecklenburg derived roster, cohorts, provenance or exclusions failed independent replay')
    return expected['coverage']


def read_zip_rows(source_id, selection, delimiter=','):
    path = ROOT / SOURCES[source_id]['path']
    member = 'Disag_2024-25_Data.txt' if source_id == 'dpi_performance' else path.stem + '.csv'
    # System unzip streams Deflate64 membership archives without cached subsets.
    process = subprocess.Popen(['unzip', '-p', str(path), member], stdout=subprocess.PIPE, text=True, encoding='utf-8-sig')
    try:
        reader = csv.DictReader(process.stdout, delimiter=delimiter)
        records = [dict(source_row=n, raw=r) for n, r in enumerate(reader, 2) if selection(r)]
        if process.wait() != 0:
            raise ValueError('Failed original source stream: ' + source_id)
        header = reader.fieldnames
    finally:
        process.stdout.close()
        if process.poll() is None:
            process.terminate(); process.wait()
    return dict(source_id=source_id, member=member, header_source_row=1, header=header, records=records)


def extract_raw_inputs():
    for metadata in SOURCES.values():
        path = ROOT / metadata['path']
        if path.stat().st_size != metadata['bytes'] or hashlib.sha256(path.read_bytes()).hexdigest() != metadata['sha256']:
            raise ValueError('Pinned original source fingerprint changed: ' + metadata['path'])
    district = lambda r: r['LEAID'] == LEA or r['ST_LEAID'] == NATIVE_LEA
    inputs = {name: read_zip_rows(source_id, district) for name, source_id in [('directory', 'ccd_directory'), ('lea_directory', 'ccd_lea_directory')]}
    inputs['membership'] = read_zip_rows('ccd_membership', lambda r: district(r) and r['TOTAL_INDICATOR'] in {'Education Unit Total', 'Subtotal 4 - By Grade'})
    inputs['lea_membership'] = read_zip_rows('ccd_lea_membership', lambda r: district(r) and r['TOTAL_INDICATOR'] in {'Education Unit Total','Subtotal 4 - By Grade'})
    inputs['assessment'] = read_zip_rows('dpi_performance', lambda r: r['school_code'].startswith(PSU) and r['subgroup'] == 'ALL', delimiter='\t')
    inputs['assessment']['selection'] = 'All native exact-PSU All Students subjects, grade totals and assessment types; only exact individual-school RG/GS/ALL MA/RD rows may enter prospective grade-school evidence.'
    workbook = openpyxl.load_workbook(ROOT / SOURCES['dpi_income']['path'], read_only=True, data_only=True)
    try:
        rows = workbook['APR 2025'].values; header = list(next(rows))
        records = [dict(source_row=n, raw=r) for n, values in enumerate(rows, 2) if (r := dict(zip(header, values)))['psu_code'] == PSU or r['school_code'].startswith(PSU)]
    finally:
        workbook.close()
    inputs['income'] = dict(source_id='dpi_income', worksheet='APR 2025', header_source_row=1, header=header,
                           selection='Complete APR2025 PSU600 records including district aggregate; cells retain published types and missingness.', records=records)
    with zipfile.ZipFile(ROOT/SOURCES['dpi_src']['path'])as archive:
        member='src_datasets_2425_2_of_2/rcd_location.xlsx';content=archive.read(member)
    workbook=openpyxl.load_workbook(io.BytesIO(content),read_only=True,data_only=True)
    try:
        sheet=workbook.worksheets[0];rows=sheet.values;header=list(next(rows));records=[]
        for n,values in enumerate(rows,2):
            r=dict(zip(header,values))
            if r['year']==YEAR and(str(r.get('agency_code','')).startswith(PSU)or r.get('lea_code')==PSU+'LEA'):
                records.append(dict(source_row=n,raw=r))
        inputs['native_location']=dict(source_id='dpi_src',member=member,member_sha256=hashlib.sha256(content).hexdigest(),member_bytes=len(content),worksheet=sheet.title,header_source_row=1,header=header,records=records)
    finally:
        workbook.close()
    definitions = {}
    for source_id in ['dpi_description', 'dpi_income_definition', 'dpi_technical']:
        reader = PdfReader(ROOT / SOURCES[source_id]['path'])
        definitions[source_id] = dict(source_id=source_id, pages=[dict(page=i, text=p.extract_text()) for i, p in enumerate(reader.pages, 1)])
    workbook=openpyxl.load_workbook(ROOT/SOURCES['dpi_src_dictionary']['path'],read_only=True,data_only=True)
    try:
        definitions['dpi_src_dictionary']=dict(source_id='dpi_src_dictionary',worksheets=[dict(worksheet=sheet.title,rows=[dict(source_row=n,values=list(values))for n,values in enumerate(sheet.values,1)])for sheet in workbook])
    finally:
        workbook.close()
    definitions['dpi_eds_web_definition'] = dict(source_id='dpi_eds_web_definition', raw_markup=(ROOT / SOURCES['dpi_eds_web_definition']['path']).read_text())
    with zipfile.ZipFile(ROOT / SOURCES['dpi_performance']['path']) as z:
        embedded = z.read('Disag_2024-25_Description.pdf')
    if hashlib.sha256(embedded).hexdigest() != SOURCES['dpi_description']['sha256']:
        raise ValueError('Embedded native performance definition differs from pinned PDF')
    definitions['embedded_description'] = dict(source_id='dpi_performance', member='Disag_2024-25_Description.pdf', sha256=hashlib.sha256(embedded).hexdigest())
    inputs['definitions'] = definitions
    payload = json.loads((ROOT / PLANNING['source']).read_text())
    record = next(r for r in payload['agencies'] if r['nces_lea_id'] == LEA)
    fields = ['nces_lea_id', 'state', 'native_lea_id', 'enrollment', 'enrollment_flag', 'school_directory_operational_count',
              'potential_cohorts', 'potential_school_ids', 'directory_raw', 'directory_source_row', 'membership_raw', 'membership_source_row']
    inputs['planning'] = dict(source=PLANNING['source'], sources=payload['sources'], record={k: record[k] for k in fields})
    return inputs


def extract(output=OUTPUT):
    snapshot = load_north_carolina(); inputs = extract_raw_inputs()
    document = dict(schema_version=1, state='NC', year=YEAR, academic_year=ACADEMIC_YEAR, nces_lea_id=LEA, native_lea_id=NATIVE_LEA,
                    status='audit_pending', scope='source_cohort_audit_only', approved_for_source=False, approved_for_modeling=False,
                    sources=SOURCES, policy=POLICY, planning=PLANNING, raw_inputs=inputs,
                    state_extract=dict(path='data/source/north-carolina.json', sha256=STATE_EXTRACT_SHA))
    document.update(derive(document, snapshot)); validate(document, snapshot)
    Path(output).write_text(json.dumps(document, separators=(',', ':'), ensure_ascii=False, allow_nan=False) + '\n')
    return document['coverage']


if __name__ == '__main__':
    parser = argparse.ArgumentParser(description=__doc__)
    parser.add_argument('--extract', action='store_true')
    parser.add_argument('--output', type=Path, default=OUTPUT)
    args = parser.parse_args()
    print(json.dumps(extract(args.output) if args.extract else validate(json.loads(args.output.read_text())), sort_keys=True))
